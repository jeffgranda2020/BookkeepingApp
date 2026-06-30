import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import csv
import os
import sys
import re
import shutil
from datetime import datetime
from PIL import Image as PILImage, ImageTk
import pgeocode
from tkcalendar import Calendar

import pdfplumber
import openpyxl
from ofxparse import OfxParser

import database as db
import autocategorize
import export_reports
import updater
import auth_ui
from version import APP_NAME, APP_VERSION, APP_AUTHOR

NOMI = pgeocode.Nominatim("us")
if getattr(sys, 'frozen', False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))


class BookkeepingApp:
    def __init__(self, root, user_id):
        self.root = root
        self.user_id = user_id
        self.root.title(f"{APP_NAME} v{APP_VERSION}")
        self.root.state("zoomed")
        self.root.minsize(900, 600)

        # Set app icon
        icon_path = os.path.join(APP_DIR, "app_icon.ico")
        if os.path.exists(icon_path):
            self.root.iconbitmap(icon_path)

        # Menu bar
        self._build_menu_bar()

        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=(10, 2))

        self._build_company_info_tab()
        self._build_accounting_tab()
        self._build_contractors_tab()
        self._build_invoicing_tab()
        self._build_sub_payments_tab()

        # Credit footer
        footer = ttk.Label(self.root, text=f"Created by {APP_AUTHOR}", font=("Segoe UI", 8), foreground="#666666")
        footer.pack(side="bottom", pady=(2, 6))

        # Check for updates in background
        self.root.after(2000, self._check_updates)

    def _center_dialog(self, dialog):
        dialog.update_idletasks()
        w = dialog.winfo_width()
        h = dialog.winfo_height()
        x = (dialog.winfo_screenwidth() // 2) - (w // 2)
        y = (dialog.winfo_screenheight() // 2) - (h // 2)
        dialog.geometry(f"+{x}+{y}")

    # ──────────────────────────────────────────────
    # Menu Bar
    # ──────────────────────────────────────────────
    def _build_menu_bar(self):
        menubar = tk.Menu(self.root)
        self.root.config(menu=menubar)

        account_menu = tk.Menu(menubar, tearoff=0)
        account_menu.add_command(label="Log Out", command=self._logout)
        menubar.add_cascade(label="Account", menu=account_menu)

        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label="Check for Updates", command=self._manual_update_check)
        help_menu.add_command(label="Update Log", command=self._show_update_log)
        help_menu.add_separator()
        help_menu.add_command(label="About", command=self._show_about)
        menubar.add_cascade(label="Help", menu=help_menu)

    def _logout(self):
        self.root.destroy()
        login = auth_ui.LoginWindow()
        if login.user_id is None:
            sys.exit(0)
        root = tk.Tk()
        BookkeepingApp(root, login.user_id)
        root.mainloop()

    def _check_updates(self):
        def on_result(new_version, changelog, download_url):
            if new_version:
                self.root.after(0, lambda: self._show_update_prompt(new_version, changelog, download_url))

        updater.check_for_updates(on_result)

    def _manual_update_check(self):
        def on_result(new_version, changelog, download_url):
            if new_version:
                self.root.after(0, lambda: self._show_update_prompt(new_version, changelog, download_url))
            else:
                self.root.after(0, lambda: messagebox.showinfo("Up to Date", f"You are running the latest version ({APP_VERSION})."))

        updater.check_for_updates(on_result)

    def _show_update_prompt(self, new_version, changelog, download_url):
        dialog = tk.Toplevel(self.root)
        dialog.title("Update Available")
        dialog.geometry("520x400")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        ttk.Label(dialog, text="A new version is available!", font=("Segoe UI", 12, "bold")).pack(pady=(15, 5))
        ttk.Label(dialog, text=f"Current: v{APP_VERSION}  →  New: v{new_version}", font=("Segoe UI", 10)).pack(pady=5)

        ttk.Label(dialog, text="What's New:", font=("Segoe UI", 10, "bold")).pack(anchor="w", padx=20, pady=(10, 5))

        text_frame = ttk.Frame(dialog)
        text_frame.pack(fill="both", expand=True, padx=20, pady=5)
        text_widget = tk.Text(text_frame, wrap="word", height=10, font=("Segoe UI", 9))
        text_scroll = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        text_widget.configure(yscrollcommand=text_scroll.set)
        text_widget.pack(side="left", fill="both", expand=True)
        text_scroll.pack(side="right", fill="y")
        text_widget.insert("1.0", changelog or "See release notes for details.")
        text_widget.config(state="disabled")

        # Progress bar (hidden until download starts)
        progress_frame = ttk.Frame(dialog)
        progress_frame.pack(fill="x", padx=20, pady=(5, 0))
        progress_var = tk.DoubleVar(value=0)
        progress_bar = ttk.Progressbar(progress_frame, variable=progress_var, maximum=100)
        progress_label = ttk.Label(progress_frame, text="", font=("Segoe UI", 8))

        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(fill="x", padx=20, pady=15)

        def install_update():
            if not download_url or not download_url.endswith(".exe"):
                import webbrowser
                webbrowser.open(download_url)
                dialog.destroy()
                return

            install_btn.config(state="disabled", text="Downloading...")
            later_btn.config(state="disabled")
            progress_bar.pack(fill="x", pady=(0, 3))
            progress_label.pack(anchor="w")

            def do_download():
                def on_progress(downloaded, total):
                    if total > 0:
                        pct = (downloaded / total) * 100
                        self.root.after(0, lambda: progress_var.set(pct))
                        mb_down = downloaded / (1024 * 1024)
                        mb_total = total / (1024 * 1024)
                        self.root.after(0, lambda: progress_label.config(text=f"{mb_down:.1f} MB / {mb_total:.1f} MB"))

                try:
                    exe_path = updater.download_update(download_url, progress_callback=on_progress)
                    self.root.after(0, lambda: self._finalize_update(exe_path, dialog))
                except Exception as e:
                    self.root.after(0, lambda: self._update_failed(str(e), dialog, install_btn, later_btn))

            import threading
            threading.Thread(target=do_download, daemon=True).start()

        install_btn = ttk.Button(btn_frame, text="Install Update", command=install_update)
        install_btn.pack(side="left", padx=5)
        later_btn = ttk.Button(btn_frame, text="Later", command=dialog.destroy)
        later_btn.pack(side="right", padx=5)

    def _finalize_update(self, exe_path, dialog):
        dialog.destroy()
        if messagebox.askyesno("Ready to Update",
                               "Download complete. The app will close and restart with the new version.\n\nContinue?"):
            updater.apply_update(exe_path)
        else:
            try:
                os.remove(exe_path)
            except OSError:
                pass

    def _update_failed(self, error_msg, dialog, install_btn, later_btn):
        messagebox.showerror("Update Failed", f"Could not download the update:\n{error_msg}", parent=dialog)
        install_btn.config(state="normal", text="Install Update")
        later_btn.config(state="normal")

    def _show_update_log(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Update Log")
        dialog.geometry("500x400")
        dialog.transient(self.root)
        dialog.resizable(True, True)

        ttk.Label(dialog, text="Update Log", font=("Segoe UI", 12, "bold")).pack(pady=(15, 10))

        text_frame = ttk.Frame(dialog)
        text_frame.pack(fill="both", expand=True, padx=15, pady=(0, 10))
        text_widget = tk.Text(text_frame, wrap="word", font=("Segoe UI", 9))
        text_scroll = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        text_widget.configure(yscrollcommand=text_scroll.set)
        text_widget.pack(side="left", fill="both", expand=True)
        text_scroll.pack(side="right", fill="y")
        text_widget.insert("1.0", updater.get_changelog_text())
        text_widget.config(state="disabled")

        ttk.Button(dialog, text="Close", command=dialog.destroy).pack(pady=10)

    def _show_about(self):
        messagebox.showinfo(
            "About",
            f"{APP_NAME}\n"
            f"Version {APP_VERSION}\n\n"
            f"Created by {APP_AUTHOR}\n\n"
            f"Bookkeeping for the trades.\n"
            f"Track invoices, manage contractors, and generate\n"
            f"professional financial statements."
        )

    # ──────────────────────────────────────────────
    # Company Information Tab
    # ──────────────────────────────────────────────
    def _build_company_info_tab(self):
        frame = ttk.Frame(self.notebook, padding=20)
        self.notebook.add(frame, text="Company Information")

        content = ttk.Frame(frame)
        content.pack(anchor="w")

        ttk.Label(content, text="Company Information", font=("Segoe UI", 16, "bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 20)
        )

        labels = [
            "Company Name",
            "EIN Number",
            "HIC Number",
            "Phone Number",
            "Email",
        ]

        self.company_entries = {}
        for i, label in enumerate(labels, start=1):
            ttk.Label(content, text=label + ":").grid(row=i, column=0, sticky="e", padx=(0, 10), pady=6)
            entry = ttk.Entry(content, width=40)
            entry.grid(row=i, column=1, sticky="w", pady=6)
            self.company_entries[label] = entry

        next_row = len(labels) + 1

        # Logo section
        logo_frame = ttk.LabelFrame(content, text="Company Logo", padding=10)
        logo_frame.grid(row=next_row, column=0, columnspan=2, sticky="w", pady=10)

        self.logo_preview_label = ttk.Label(logo_frame, text="No logo uploaded")
        self.logo_preview_label.pack(side="left", padx=(0, 15))

        ttk.Button(logo_frame, text="Upload Logo", command=self._upload_logo).pack(side="left", padx=5)
        ttk.Button(logo_frame, text="Remove Logo", command=self._remove_logo).pack(side="left", padx=5)

        btn_frame = ttk.Frame(content)
        btn_frame.grid(row=next_row + 1, column=0, columnspan=2, sticky="w", pady=20)
        ttk.Button(btn_frame, text="Save", command=self._save_company_info).pack(side="left", padx=5)

        self._load_company_info()
        self._update_logo_preview()

    def _save_company_info(self):
        data = {k: v.get() for k, v in self.company_entries.items()}
        db.save_company_info(data, self.user_id)
        messagebox.showinfo("Saved", "Company information saved.")

    def _load_company_info(self):
        data = db.load_company_info(self.user_id)
        for key, entry in self.company_entries.items():
            if key in data:
                entry.insert(0, data[key])

    def _get_logo_path(self):
        return os.path.join(APP_DIR, "logo.png")

    def _upload_logo(self):
        filepath = filedialog.askopenfilename(
            title="Select Company Logo",
            filetypes=[
                ("Image Files", "*.png;*.jpg;*.jpeg;*.bmp;*.gif;*.tiff"),
                ("All Files", "*.*"),
            ],
        )
        if not filepath:
            return

        logo_dest = self._get_logo_path()
        img = PILImage.open(filepath)
        img = img.convert("RGBA")
        img.save(logo_dest, "PNG")
        self._update_logo_preview()
        messagebox.showinfo("Logo Uploaded", "Company logo saved successfully.")

    def _remove_logo(self):
        logo_path = self._get_logo_path()
        if os.path.exists(logo_path):
            os.remove(logo_path)
            self._update_logo_preview()
            messagebox.showinfo("Removed", "Company logo removed.")

    def _update_logo_preview(self):
        logo_path = self._get_logo_path()
        if os.path.exists(logo_path):
            img = PILImage.open(logo_path)
            img.thumbnail((80, 80))
            self._logo_photo = ImageTk.PhotoImage(img)
            self.logo_preview_label.config(image=self._logo_photo, text="")
        else:
            self._logo_photo = None
            self.logo_preview_label.config(image="", text="No logo uploaded")

    # ──────────────────────────────────────────────
    # Sorting & Filtering Helpers (Excel-style)
    # ──────────────────────────────────────────────
    def _setup_sortable(self, tree, refresh_func):
        tree._sort_state = {}
        tree._col_filters = {}
        tree._refresh_func = refresh_func

        def on_heading_click(col):
            reverse = tree._sort_state.get(col, False)
            tree._sort_state[col] = not reverse
            items = [(tree.set(k, col), k) for k in tree.get_children("")]
            try:
                items.sort(key=lambda x: float(x[0].replace("$", "").replace(",", "")), reverse=reverse)
            except ValueError:
                items.sort(key=lambda x: x[0].lower(), reverse=reverse)
            for index, (val, k) in enumerate(items):
                tree.move(k, "", index)

        def on_right_click(event):
            region = tree.identify_region(event.x, event.y)
            if region == "heading":
                col_id = tree.identify_column(event.x)
                col_idx = int(col_id.replace("#", "")) - 1
                col = tree["columns"][col_idx]
                self._show_column_filter(tree, col, event)

        for col in tree["columns"]:
            tree.heading(col, command=lambda c=col: on_heading_click(c))
        tree.bind("<Button-3>", on_right_click)

    def _show_column_filter(self, tree, col, event):
        all_values = set()
        for item in tree.get_children(""):
            all_values.add(tree.set(item, col))

        menu = tk.Menu(self.root, tearoff=0)
        menu.add_command(label="Sort A → Z", command=lambda: self._apply_sort(tree, col, False))
        menu.add_command(label="Sort Z → A", command=lambda: self._apply_sort(tree, col, True))
        menu.add_separator()
        menu.add_command(label="Clear Filter", command=lambda: self._clear_col_filter(tree, col))
        menu.add_separator()

        sorted_vals = sorted(all_values, key=str.lower)
        for val in sorted_vals[:20]:
            display = val if val else "(empty)"
            menu.add_command(label=f"Show: {display}", command=lambda v=val: self._apply_col_filter(tree, col, v))

        menu.tk_popup(event.x_root, event.y_root)

    def _apply_sort(self, tree, col, reverse):
        items = [(tree.set(k, col), k) for k in tree.get_children("")]
        try:
            items.sort(key=lambda x: float(x[0].replace("$", "").replace(",", "")), reverse=reverse)
        except ValueError:
            items.sort(key=lambda x: x[0].lower(), reverse=reverse)
        for index, (val, k) in enumerate(items):
            tree.move(k, "", index)

    def _apply_col_filter(self, tree, col, value):
        tree._col_filters[col] = value
        tree._refresh_func()

    def _clear_col_filter(self, tree, col):
        tree._col_filters.pop(col, None)
        tree._refresh_func()

    def _passes_filters(self, tree, values, columns):
        if not hasattr(tree, "_col_filters"):
            return True
        for col, filter_val in tree._col_filters.items():
            col_idx = list(columns).index(col)
            if values[col_idx] != filter_val:
                return False
        return True

    # ──────────────────────────────────────────────
    # Phone Format Helper
    # ──────────────────────────────────────────────
    def _format_phone_display(self, phone):
        digits = re.sub(r"\D", "", phone or "")
        if len(digits) == 10:
            return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
        return phone or ""

    def _format_phone_entry(self, entry):
        def on_key(*args):
            text = entry.get()
            digits = re.sub(r"\D", "", text)
            if len(digits) > 10:
                digits = digits[:10]
            if len(digits) >= 7:
                formatted = f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
            elif len(digits) >= 4:
                formatted = f"({digits[:3]}) {digits[3:]}"
            elif len(digits) >= 1:
                formatted = f"({digits}"
            else:
                formatted = ""
            if formatted != text:
                entry.delete(0, "end")
                entry.insert(0, formatted)
                entry.icursor("end")

        entry.bind("<KeyRelease>", on_key)

    # ──────────────────────────────────────────────
    # Date Picker Helper
    # ──────────────────────────────────────────────
    def _create_date_entry(self, parent, default=None, min_date=None, max_date=None):
        frame = ttk.Frame(parent)
        var = tk.StringVar(value=default or datetime.now().strftime("%Y-%m-%d"))
        entry = ttk.Entry(frame, textvariable=var, width=12)
        entry.pack(side="left")

        def pick_date():
            top = tk.Toplevel(parent)
            top.title("Select Date")
            top.geometry("300x250")
            top.transient(parent)
            top.grab_set()
            top.resizable(False, False)
            cal_kwargs = {"selectmode": "day", "date_pattern": "yyyy-mm-dd"}
            if min_date:
                cal_kwargs["mindate"] = min_date
            if max_date:
                cal_kwargs["maxdate"] = max_date
            cal = Calendar(top, **cal_kwargs)
            cal.pack(fill="both", expand=True, padx=10, pady=10)
            def select():
                var.set(cal.get_date())
                top.destroy()
            ttk.Button(top, text="Select", command=select).pack(pady=5)

        ttk.Button(frame, text="📅", width=3, command=pick_date).pack(side="left", padx=2)
        return frame, var

    # ──────────────────────────────────────────────
    # Accounting Tab
    # ──────────────────────────────────────────────
    def _build_accounting_tab(self):
        frame = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(frame, text="Accounting")

        # Top toolbar
        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))

        ttk.Button(toolbar, text="Import Statement", command=self._import_statement).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Add Transaction", command=self._add_transaction_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Add Category", command=self._add_category_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Categorize Selected", command=self._categorize_selected).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete Selected", command=self._delete_selected_txn).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Select All", command=self._select_all_txn).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Deselect All", command=self._deselect_all_txn).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Generate P&L", command=self._generate_pl).pack(side="right", padx=5)
        ttk.Button(toolbar, text="Generate Balance Sheet", command=self._generate_balance_sheet).pack(side="right", padx=5)

        # Search bar
        search_frame = ttk.Frame(frame)
        search_frame.pack(fill="x", pady=(0, 5))
        ttk.Label(search_frame, text="Search:").pack(side="left", padx=(0, 5))
        self._txn_filter_var = tk.StringVar()
        ttk.Entry(search_frame, textvariable=self._txn_filter_var, width=30).pack(side="left", padx=(0, 5))
        ttk.Button(search_frame, text="Clear", command=lambda: self._txn_filter_var.set("")).pack(side="left", padx=(0, 10))
        ttk.Label(search_frame, text="(Right-click column headers to filter/sort)", font=("Segoe UI", 8)).pack(side="left")
        self._txn_filter_var.trace_add("write", lambda *a: self._refresh_transactions())

        # Transactions treeview
        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("Date", "Description", "Amount", "Category Type", "Category", "Source")
        self.txn_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=18, selectmode="extended")
        for col in cols:
            self.txn_tree.heading(col, text=col, anchor="center")
            width = 100 if col in ("Date", "Amount", "Source") else 160
            anchor = "center"
            self.txn_tree.column(col, width=width, anchor=anchor, stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.txn_tree.yview)
        self.txn_tree.configure(yscrollcommand=scrollbar.set)
        self.txn_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.txn_tree.bind("<Double-1>", self._edit_transaction_category)
        self._setup_sortable(self.txn_tree, self._refresh_transactions)

        self._refresh_transactions()

    def _refresh_transactions(self):
        for row in self.txn_tree.get_children():
            self.txn_tree.delete(row)
        query = self._txn_filter_var.get().lower() if hasattr(self, "_txn_filter_var") else ""
        cols = self.txn_tree["columns"]
        transactions = db.get_transactions(user_id=self.user_id)
        for txn in transactions:
            tid, date, desc, amount, cat_type, cat_name, source = txn
            values = (date, desc or "", f"${amount:,.2f}", cat_type or "", cat_name or "Uncategorized", source or "")
            if query and not any(query in str(v).lower() for v in values):
                continue
            if not self._passes_filters(self.txn_tree, values, cols):
                continue
            self.txn_tree.insert("", "end", iid=tid, values=values)

    def _select_all_txn(self):
        items = self.txn_tree.get_children()
        self.txn_tree.selection_set(items)

    def _deselect_all_txn(self):
        self.txn_tree.selection_remove(self.txn_tree.selection())

    def _delete_selected_txn(self):
        selected = self.txn_tree.selection()
        if not selected:
            messagebox.showwarning("No Selection", "Select one or more transactions to delete.")
            return
        count = len(selected)
        if not messagebox.askyesno("Confirm", f"Delete {count} transaction(s)? This cannot be undone."):
            return
        txn_ids = [int(item) for item in selected]
        db.delete_transactions(txn_ids)
        self._refresh_transactions()

    def _add_transaction_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Add Transaction")
        dialog.geometry("450x320")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 12, "pady": 8}

        form_frame = ttk.LabelFrame(dialog, text="Transaction Details", padding=10)
        form_frame.pack(fill="x", padx=15, pady=(15, 5))

        ttk.Label(form_frame, text="Date:").grid(row=0, column=0, sticky="e", **pad)
        date_entry = ttk.Entry(form_frame, width=28)
        date_entry.insert(0, datetime.now().strftime("%Y-%m-%d"))
        date_entry.grid(row=0, column=1, sticky="w", **pad)

        ttk.Label(form_frame, text="Description:").grid(row=1, column=0, sticky="e", **pad)
        desc_entry = ttk.Entry(form_frame, width=28)
        desc_entry.grid(row=1, column=1, sticky="w", **pad)

        ttk.Label(form_frame, text="Amount:").grid(row=2, column=0, sticky="e", **pad)
        amount_entry = ttk.Entry(form_frame, width=28)
        amount_entry.grid(row=2, column=1, sticky="w", **pad)

        ttk.Label(form_frame, text="Category:").grid(row=3, column=0, sticky="e", **pad)
        categories = db.get_categories(user_id=self.user_id)
        cat_display = [f"{c[1]} - {c[2]}" for c in categories]
        cat_ids = [c[0] for c in categories]
        cat_var = tk.StringVar()
        cat_combo = ttk.Combobox(form_frame, textvariable=cat_var, values=cat_display, width=26, state="readonly")
        cat_combo.grid(row=3, column=1, sticky="w", **pad)

        def save():
            try:
                amount = float(amount_entry.get().replace(",", "").replace("$", ""))
            except ValueError:
                messagebox.showerror("Error", "Invalid amount.")
                return
            date_val = date_entry.get().strip()
            cat_idx = cat_combo.current()
            cat_id = cat_ids[cat_idx] if cat_idx >= 0 else None
            db.add_transaction(date_val, desc_entry.get().strip(), amount, cat_id, user_id=self.user_id)
            dialog.destroy()
            self._refresh_transactions()

        ttk.Button(dialog, text="Save", command=save).pack(pady=15)

    def _edit_transaction_category(self, event):
        item = self.txn_tree.focus()
        if not item:
            return
        txn_id = int(item)

        dialog = tk.Toplevel(self.root)
        dialog.title("Assign Category")
        dialog.geometry("350x180")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)

        ttk.Label(dialog, text="Category:").pack(pady=10)
        categories = db.get_categories(user_id=self.user_id)
        cat_display = [f"{c[1]} - {c[2]}" for c in categories]
        cat_ids = [c[0] for c in categories]
        cat_var = tk.StringVar()
        cat_combo = ttk.Combobox(dialog, textvariable=cat_var, values=cat_display, width=35, state="readonly")
        cat_combo.pack(pady=5)

        def save():
            cat_idx = cat_combo.current()
            if cat_idx >= 0:
                db.update_transaction_category(txn_id, cat_ids[cat_idx])
            dialog.destroy()
            self._refresh_transactions()

        ttk.Button(dialog, text="Save", command=save).pack(pady=10)

    def _categorize_selected(self):
        selected = self.txn_tree.selection()
        if not selected:
            messagebox.showwarning("No Selection", "Select one or more transactions to categorize.\n\nTip: Hold Ctrl or Shift to select multiple rows.")
            return

        txn_ids = [int(item) for item in selected]

        dialog = tk.Toplevel(self.root)
        dialog.title(f"Categorize {len(txn_ids)} Transaction(s)")
        dialog.geometry("400x180")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)

        ttk.Label(dialog, text=f"Assign category to {len(txn_ids)} selected transaction(s):").pack(pady=10)
        categories = db.get_categories(user_id=self.user_id)
        cat_display = [f"{c[1]} - {c[2]}" for c in categories]
        cat_ids = [c[0] for c in categories]
        cat_var = tk.StringVar()
        cat_combo = ttk.Combobox(dialog, textvariable=cat_var, values=cat_display, width=35, state="readonly")
        cat_combo.pack(pady=5)

        def save():
            cat_idx = cat_combo.current()
            if cat_idx >= 0:
                for txn_id in txn_ids:
                    db.update_transaction_category(txn_id, cat_ids[cat_idx])
            dialog.destroy()
            self._refresh_transactions()

        ttk.Button(dialog, text="Apply", command=save).pack(pady=10)

    def _add_category_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Add Custom Category")
        dialog.geometry("420x230")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 12, "pady": 8}

        form_frame = ttk.LabelFrame(dialog, text="Category Details", padding=10)
        form_frame.pack(fill="x", padx=15, pady=(15, 5))

        ttk.Label(form_frame, text="Type:").grid(row=0, column=0, sticky="e", **pad)
        type_var = tk.StringVar()
        type_combo = ttk.Combobox(
            form_frame, textvariable=type_var,
            values=["Income", "COGS", "Expense", "Asset", "Liability", "Equity"],
            width=26, state="readonly"
        )
        type_combo.grid(row=0, column=1, sticky="w", **pad)
        type_combo.current(2)

        ttk.Label(form_frame, text="Name:").grid(row=1, column=0, sticky="e", **pad)
        name_entry = ttk.Entry(form_frame, width=28)
        name_entry.grid(row=1, column=1, sticky="w", **pad)

        def save():
            cat_type = type_var.get()
            cat_name = name_entry.get().strip()
            if not cat_name:
                messagebox.showerror("Error", "Enter a category name.")
                return
            added = db.add_category(cat_type, cat_name, user_id=self.user_id)
            if added:
                messagebox.showinfo("Added", f"Category '{cat_name}' added under {cat_type}.")
            else:
                messagebox.showwarning("Exists", f"Category '{cat_name}' already exists under {cat_type}.")
            dialog.destroy()

        ttk.Button(dialog, text="Save", command=save).pack(pady=15)

    def _import_statement(self):
        filepath = filedialog.askopenfilename(
            title="Select Statement File",
            filetypes=[
                ("All Supported", "*.csv;*.xlsx;*.xls;*.pdf;*.ofx;*.qfx"),
                ("CSV Files", "*.csv"),
                ("Excel Files", "*.xlsx;*.xls"),
                ("PDF Files", "*.pdf"),
                ("OFX/QFX Files", "*.ofx;*.qfx"),
                ("All Files", "*.*"),
            ],
        )
        if not filepath:
            return

        ext = os.path.splitext(filepath)[1].lower()
        try:
            if ext == ".csv":
                rows = self._parse_csv(filepath)
            elif ext in (".xlsx", ".xls"):
                rows = self._parse_excel(filepath)
            elif ext == ".pdf":
                rows = self._parse_pdf(filepath)
            elif ext in (".ofx", ".qfx"):
                rows = self._parse_ofx(filepath)
            else:
                messagebox.showerror("Error", f"Unsupported file type: {ext}")
                return
        except Exception as e:
            messagebox.showerror("Import Error", str(e))
            return

        if not rows:
            messagebox.showwarning("Empty", "No transactions found in the file.")
            return

        self._preview_import(rows, filepath)

    def _preview_import(self, rows, filepath):
        dialog = tk.Toplevel(self.root)
        dialog.title(f"Review Import - {os.path.basename(filepath)}")
        dialog.geometry("900x520")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)

        ttk.Label(dialog, text=f"Found {len(rows)} transactions. Double-click to edit, then Import.", font=("Segoe UI", 10, "bold")).pack(pady=10)

        tree_frame = ttk.Frame(dialog)
        tree_frame.pack(fill="both", expand=True, padx=10, pady=5)

        cols = ("Date", "Description", "Amount", "Category")
        tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=18)
        tree.heading("Date", text="Date", anchor="center")
        tree.heading("Description", text="Description", anchor="center")
        tree.heading("Amount", text="Amount", anchor="center")
        tree.heading("Category", text="Category", anchor="center")
        tree.column("Date", width=90, anchor="center", stretch=True)
        tree.column("Description", width=360, anchor="center", stretch=True)
        tree.column("Amount", width=100, anchor="center", stretch=True)
        tree.column("Category", width=220, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        import_data = []
        for r in rows:
            cat_type, cat_name = autocategorize.guess_category(r[1], r[2])
            cat_label = f"{cat_type} - {cat_name}" if cat_type else "Uncategorized"
            import_data.append([r[0], r[1], r[2], cat_label])

        for i, row_data in enumerate(import_data):
            tree.insert("", "end", iid=str(i), values=(row_data[0], row_data[1], f"${row_data[2]:,.2f}", row_data[3]))

        def edit_row(event):
            item = tree.focus()
            if not item:
                return
            idx = int(item)
            col_region = tree.identify_column(event.x)
            col_idx = int(col_region.replace("#", "")) - 1

            edit_dialog = tk.Toplevel(dialog)
            edit_dialog.title("Edit Transaction")
            edit_dialog.geometry("450x250")
            edit_dialog.transient(dialog)
            edit_dialog.grab_set()
            self._center_dialog(edit_dialog)
            edit_dialog.resizable(False, False)

            row_data = import_data[idx]
            pad = {"padx": 10, "pady": 6}

            ttk.Label(edit_dialog, text="Date:").grid(row=0, column=0, sticky="e", **pad)
            date_e = ttk.Entry(edit_dialog, width=15)
            date_e.insert(0, row_data[0])
            date_e.grid(row=0, column=1, sticky="w", **pad)

            ttk.Label(edit_dialog, text="Description:").grid(row=1, column=0, sticky="e", **pad)
            desc_e = ttk.Entry(edit_dialog, width=40)
            desc_e.insert(0, row_data[1])
            desc_e.grid(row=1, column=1, sticky="w", **pad)

            ttk.Label(edit_dialog, text="Amount:").grid(row=2, column=0, sticky="e", **pad)
            amt_e = ttk.Entry(edit_dialog, width=15)
            amt_e.insert(0, f"{row_data[2]:.2f}")
            amt_e.grid(row=2, column=1, sticky="w", **pad)

            ttk.Label(edit_dialog, text="Category:").grid(row=3, column=0, sticky="e", **pad)
            categories = db.get_categories(user_id=self.user_id)
            cat_display = ["Uncategorized"] + [f"{c[1]} - {c[2]}" for c in categories]
            cat_var = tk.StringVar(value=row_data[3])
            cat_combo = ttk.Combobox(edit_dialog, textvariable=cat_var, values=cat_display, width=35, state="readonly")
            cat_combo.grid(row=3, column=1, sticky="w", **pad)
            if row_data[3] in cat_display:
                cat_combo.current(cat_display.index(row_data[3]))

            def save_edit():
                try:
                    new_amt = float(amt_e.get().replace("$", "").replace(",", ""))
                except ValueError:
                    messagebox.showerror("Error", "Invalid amount.", parent=edit_dialog)
                    return
                import_data[idx] = [date_e.get().strip(), desc_e.get().strip(), new_amt, cat_var.get()]
                tree.item(item, values=(import_data[idx][0], import_data[idx][1], f"${new_amt:,.2f}", import_data[idx][3]))
                edit_dialog.destroy()

            ttk.Button(edit_dialog, text="Save", command=save_edit).grid(row=4, column=0, columnspan=2, pady=15)

            if col_idx == 0:
                date_e.focus_set()
            elif col_idx == 1:
                desc_e.focus_set()
            elif col_idx == 2:
                amt_e.focus_set()
            elif col_idx == 3:
                cat_combo.focus_set()

        def delete_selected():
            selected = tree.selection()
            if not selected:
                return
            for item in sorted(selected, key=lambda x: int(x), reverse=True):
                idx = int(item)
                import_data[idx] = None
                tree.delete(item)
            while None in import_data:
                import_data.remove(None)
            for item in tree.get_children():
                tree.delete(item)
            for i, row_data in enumerate(import_data):
                tree.insert("", "end", iid=str(i), values=(row_data[0], row_data[1], f"${row_data[2]:,.2f}", row_data[3]))

        tree.bind("<Double-1>", edit_row)

        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(fill="x", pady=10, padx=10)

        def confirm():
            source = os.path.basename(filepath)
            categories = db.get_categories(user_id=self.user_id)
            cat_lookup = {f"{c[1]} - {c[2]}": c[0] for c in categories}
            db_rows = []
            for row_data in import_data:
                cat_id = cat_lookup.get(row_data[3])
                db_rows.append((row_data[0], row_data[1], row_data[2], cat_id, source))
            db.add_transactions_bulk(db_rows, user_id=self.user_id)
            self._refresh_transactions()
            dialog.destroy()
            messagebox.showinfo("Import Complete", f"Imported {len(db_rows)} transactions.")

        ttk.Button(btn_frame, text="Delete Selected", command=delete_selected).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Import", command=confirm).pack(side="right", padx=5)
        ttk.Button(btn_frame, text="Cancel", command=dialog.destroy).pack(side="right", padx=5)

    # ── Parsers ──

    def _parse_csv(self, filepath):
        with open(filepath, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            header = [h.strip().lower() for h in next(reader)]

            date_col = self._find_col(header, ["date", "transaction date", "post date", "posted date"])
            desc_col = self._find_col(header, ["description", "memo", "name", "payee", "transaction description"])
            amount_col = self._find_col(header, ["amount", "total", "transaction amount"])
            debit_col = self._find_col(header, ["debit", "withdrawal", "charge"])
            credit_col = self._find_col(header, ["credit", "deposit", "payment"])

            if date_col is None or desc_col is None:
                raise ValueError("Could not identify Date and Description columns in CSV.")

            rows = []
            for line in reader:
                if not line or all(c.strip() == "" for c in line):
                    continue
                date_val = line[date_col].strip()
                desc_val = line[desc_col].strip()

                if amount_col is not None:
                    raw = line[amount_col].strip().replace(",", "").replace("$", "").replace('"', "")
                    try:
                        amount_val = float(raw)
                    except ValueError:
                        continue
                elif debit_col is not None or credit_col is not None:
                    amount_val = 0.0
                    if debit_col is not None and line[debit_col].strip():
                        raw = line[debit_col].strip().replace(",", "").replace("$", "").replace('"', "")
                        try:
                            amount_val = -abs(float(raw))
                        except ValueError:
                            pass
                    if credit_col is not None and line[credit_col].strip():
                        raw = line[credit_col].strip().replace(",", "").replace("$", "").replace('"', "")
                        try:
                            amount_val = abs(float(raw))
                        except ValueError:
                            pass
                else:
                    raise ValueError("Could not identify Amount column in CSV.")

                rows.append((date_val, desc_val, amount_val))
            return rows

    def _parse_excel(self, filepath):
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not all_rows:
            return []

        header = [str(h).strip().lower() if h else "" for h in all_rows[0]]
        date_col = self._find_col(header, ["date", "transaction date", "post date", "posted date"])
        desc_col = self._find_col(header, ["description", "memo", "name", "payee", "transaction description"])
        amount_col = self._find_col(header, ["amount", "total", "transaction amount"])
        debit_col = self._find_col(header, ["debit", "withdrawal", "charge"])
        credit_col = self._find_col(header, ["credit", "deposit", "payment"])

        if date_col is None or desc_col is None:
            raise ValueError("Could not identify Date and Description columns in Excel file.")

        rows = []
        for line in all_rows[1:]:
            if not line or all(c is None or str(c).strip() == "" for c in line):
                continue
            date_val = str(line[date_col]).strip() if line[date_col] else ""
            desc_val = str(line[desc_col]).strip() if line[desc_col] else ""

            if amount_col is not None and line[amount_col] is not None:
                raw = str(line[amount_col]).replace(",", "").replace("$", "")
                try:
                    amount_val = float(raw)
                except ValueError:
                    continue
            elif debit_col is not None or credit_col is not None:
                amount_val = 0.0
                if debit_col is not None and line[debit_col]:
                    raw = str(line[debit_col]).replace(",", "").replace("$", "")
                    try:
                        amount_val = -abs(float(raw))
                    except ValueError:
                        pass
                if credit_col is not None and line[credit_col]:
                    raw = str(line[credit_col]).replace(",", "").replace("$", "")
                    try:
                        amount_val = abs(float(raw))
                    except ValueError:
                        pass
            else:
                raise ValueError("Could not identify Amount column in Excel file.")

            if date_val:
                rows.append((date_val, desc_val, amount_val))
        return rows

    def _parse_pdf(self, filepath):
        rows = []
        with pdfplumber.open(filepath) as pdf:
            # Try structured tables first
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    if not table:
                        continue
                    header = [str(h).strip().lower() if h else "" for h in table[0]]
                    date_col = self._find_col(header, ["date", "transaction date", "post date", "posted date"])
                    desc_col = self._find_col(header, ["description", "memo", "name", "payee", "transaction description"])
                    amount_col = self._find_col(header, ["amount", "total", "transaction amount"])
                    debit_col = self._find_col(header, ["debit", "withdrawal", "charge"])
                    credit_col = self._find_col(header, ["credit", "deposit", "payment"])

                    if date_col is None or desc_col is None:
                        continue

                    for line in table[1:]:
                        if not line or all(c is None or str(c).strip() == "" for c in line):
                            continue
                        date_val = str(line[date_col]).strip() if line[date_col] else ""
                        desc_val = str(line[desc_col]).strip() if line[desc_col] else ""

                        if amount_col is not None and line[amount_col]:
                            raw = str(line[amount_col]).replace(",", "").replace("$", "").replace("(", "-").replace(")", "")
                            try:
                                amount_val = float(raw)
                            except ValueError:
                                continue
                        elif debit_col is not None or credit_col is not None:
                            amount_val = 0.0
                            if debit_col is not None and line[debit_col]:
                                raw = str(line[debit_col]).replace(",", "").replace("$", "").replace("(", "-").replace(")", "")
                                try:
                                    amount_val = -abs(float(raw))
                                except ValueError:
                                    pass
                            if credit_col is not None and line[credit_col]:
                                raw = str(line[credit_col]).replace(",", "").replace("$", "").replace("(", "-").replace(")", "")
                                try:
                                    amount_val = abs(float(raw))
                                except ValueError:
                                    pass
                        else:
                            continue

                        if date_val:
                            rows.append((date_val, desc_val, amount_val))

            # Fallback: text-based parsing for unstructured PDFs (BofA, Chase, etc.)
            if not rows:
                rows = self._parse_pdf_text(pdf)

        if not rows:
            raise ValueError("Could not extract transactions from PDF. The format may not be supported.")
        return rows

    def _parse_pdf_text(self, pdf):
        # Pattern: date at start of line, amount at end (with optional negative)
        # Matches formats like: 01/15 DESCRIPTION 1,234.56  or  01/15/24 DESCRIPTION -45.00
        date_pattern = re.compile(
            r"^(\d{1,2}/\d{1,2}(?:/\d{2,4})?)\s+"  # date
            r"(.+?)\s+"                               # description
            r"(-?\$?[\d,]+\.\d{2})$"                  # amount
        )
        # Some statements put the amount on a separate column position
        date_amount_pattern = re.compile(
            r"^(\d{1,2}/\d{1,2}(?:/\d{2,4})?)\s+"  # date
            r"(.+?)\s{2,}"                            # description (followed by 2+ spaces)
            r"(-?\$?[\d,]+\.\d{2})"                   # amount
        )

        rows = []
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue
            for line in text.split("\n"):
                line = line.strip()
                if not line:
                    continue

                match = date_pattern.match(line) or date_amount_pattern.match(line)
                if match:
                    date_val = match.group(1)
                    desc_val = match.group(2).strip()
                    amount_raw = match.group(3).replace(",", "").replace("$", "")
                    try:
                        amount_val = float(amount_raw)
                    except ValueError:
                        continue

                    # Skip header-like lines and totals
                    desc_lower = desc_val.lower()
                    if any(skip in desc_lower for skip in [
                        "beginning balance", "ending balance", "total",
                        "statement period", "page", "account number"
                    ]):
                        continue

                    rows.append((date_val, desc_val, amount_val))

        return rows

    def _parse_ofx(self, filepath):
        with open(filepath, "rb") as f:
            ofx = OfxParser.parse(f)

        rows = []
        for account in ofx.accounts if hasattr(ofx, "accounts") else [ofx.account]:
            for txn in account.statement.transactions:
                date_val = txn.date.strftime("%Y-%m-%d")
                desc_val = txn.payee or txn.memo or ""
                amount_val = float(txn.amount)
                rows.append((date_val, desc_val, amount_val))
        return rows

    def _find_col(self, header, candidates):
        for i, h in enumerate(header):
            for c in candidates:
                if c in h:
                    return i
        return None

    # ──────────────────────────────────────────────
    # Contractors Tab
    # ──────────────────────────────────────────────
    def _build_contractors_tab(self):
        frame = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(frame, text="Contractors")

        # Sub-notebook for Primary / Subcontractors
        self.contractor_notebook = ttk.Notebook(frame)
        self.contractor_notebook.pack(fill="both", expand=True)

        self._build_primary_tab()
        self._build_subcontractors_tab()

    def _build_primary_tab(self):
        frame = ttk.Frame(self.contractor_notebook, padding=10)
        self.contractor_notebook.add(frame, text="Primary")

        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))
        ttk.Button(toolbar, text="Add Primary", command=self._add_primary_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete Selected", command=self._delete_primary).pack(side="left", padx=5)

        # Search
        search_frame = ttk.Frame(frame)
        search_frame.pack(fill="x", pady=(0, 5))
        ttk.Label(search_frame, text="Search:").pack(side="left", padx=(0, 5))
        self._primary_filter_var = tk.StringVar()
        ttk.Entry(search_frame, textvariable=self._primary_filter_var, width=30).pack(side="left", padx=(0, 5))
        ttk.Button(search_frame, text="Clear", command=lambda: self._primary_filter_var.set("")).pack(side="left")
        self._primary_filter_var.trace_add("write", lambda *a: self._refresh_primary())

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("ID", "Name", "Address", "City", "State", "Zip", "Phone", "Email", "EIN")
        self.primary_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=14)
        for col in cols:
            self.primary_tree.heading(col, text=col, anchor="center")
        self.primary_tree.column("ID", width=50, anchor="center", stretch=True)
        self.primary_tree.column("Name", width=150, anchor="center", stretch=True)
        self.primary_tree.column("Address", width=140, anchor="center", stretch=True)
        self.primary_tree.column("City", width=90, anchor="center", stretch=True)
        self.primary_tree.column("State", width=45, anchor="center", stretch=True)
        self.primary_tree.column("Zip", width=50, anchor="center", stretch=True)
        self.primary_tree.column("Phone", width=100, anchor="center", stretch=True)
        self.primary_tree.column("Email", width=140, anchor="center", stretch=True)
        self.primary_tree.column("EIN", width=90, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.primary_tree.yview)
        self.primary_tree.configure(yscrollcommand=scrollbar.set)
        self.primary_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self._setup_sortable(self.primary_tree, self._refresh_primary)
        self._refresh_primary()

    def _build_subcontractors_tab(self):
        frame = ttk.Frame(self.contractor_notebook, padding=10)
        self.contractor_notebook.add(frame, text="Subcontractors")

        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))
        ttk.Button(toolbar, text="Add Subcontractor", command=self._add_sub_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete Selected", command=self._delete_sub).pack(side="left", padx=5)

        # Search
        search_frame = ttk.Frame(frame)
        search_frame.pack(fill="x", pady=(0, 5))
        ttk.Label(search_frame, text="Search:").pack(side="left", padx=(0, 5))
        self._sub_filter_var = tk.StringVar()
        ttk.Entry(search_frame, textvariable=self._sub_filter_var, width=30).pack(side="left", padx=(0, 5))
        ttk.Button(search_frame, text="Clear", command=lambda: self._sub_filter_var.set("")).pack(side="left")
        self._sub_filter_var.trace_add("write", lambda *a: self._refresh_subs())

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("ID", "Name", "Address", "City", "State", "Zip", "Phone", "Email", "SSN/TIN")
        self.sub_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=14)
        for col in cols:
            self.sub_tree.heading(col, text=col, anchor="center")
        self.sub_tree.column("ID", width=50, anchor="center", stretch=True)
        self.sub_tree.column("Name", width=150, anchor="center", stretch=True)
        self.sub_tree.column("Address", width=140, anchor="center", stretch=True)
        self.sub_tree.column("City", width=90, anchor="center", stretch=True)
        self.sub_tree.column("State", width=45, anchor="center", stretch=True)
        self.sub_tree.column("Zip", width=50, anchor="center", stretch=True)
        self.sub_tree.column("Phone", width=100, anchor="center", stretch=True)
        self.sub_tree.column("Email", width=140, anchor="center", stretch=True)
        self.sub_tree.column("SSN/TIN", width=90, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.sub_tree.yview)
        self.sub_tree.configure(yscrollcommand=scrollbar.set)
        self.sub_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self._setup_sortable(self.sub_tree, self._refresh_subs)
        self._refresh_subs()

    def _refresh_primary(self):
        for row in self.primary_tree.get_children():
            self.primary_tree.delete(row)
        query = self._primary_filter_var.get().lower() if hasattr(self, "_primary_filter_var") else ""
        cols = self.primary_tree["columns"]
        contractors = db.get_contractors("Primary", user_id=self.user_id)
        for c in contractors:
            cid, code_id, name, street, city, state, zipcode, phone, email, ein, ssn_tin = c
            values = (code_id, name, street or "", city or "", state or "", zipcode or "", phone or "", email or "", ein or "")
            if query and not any(query in str(v).lower() for v in values):
                continue
            if not self._passes_filters(self.primary_tree, values, cols):
                continue
            self.primary_tree.insert("", "end", iid=cid, values=values)

    def _refresh_subs(self):
        for row in self.sub_tree.get_children():
            self.sub_tree.delete(row)
        query = self._sub_filter_var.get().lower() if hasattr(self, "_sub_filter_var") else ""
        cols = self.sub_tree["columns"]
        contractors = db.get_contractors("Subcontractor", user_id=self.user_id)
        for c in contractors:
            cid, code_id, name, street, city, state, zipcode, phone, email, ein, ssn_tin = c
            values = (code_id, name, street or "", city or "", state or "", zipcode or "", phone or "", email or "", ssn_tin or "")
            if query and not any(query in str(v).lower() for v in values):
                continue
            if not self._passes_filters(self.sub_tree, values, cols):
                continue
            self.sub_tree.insert("", "end", iid=cid, values=values)

    def _lookup_zip(self, zip_entry, city_entry, state_entry):
        zipcode = zip_entry.get().strip()
        if len(zipcode) == 5:
            result = NOMI.query_postal_code(zipcode)
            if result is not None and not (isinstance(result["place_name"], float)):
                city_entry.delete(0, "end")
                city_entry.insert(0, result["place_name"])
                state_entry.delete(0, "end")
                state_entry.insert(0, result["state_code"])

    def _add_primary_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Add Primary Contractor")
        dialog.geometry("480x420")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 12, "pady": 6}
        entries = {}

        # Info section
        info_frame = ttk.LabelFrame(dialog, text="Contractor Info", padding=10)
        info_frame.pack(fill="x", padx=15, pady=(15, 5))

        ttk.Label(info_frame, text="Name:").grid(row=0, column=0, sticky="e", **pad)
        entries["Name"] = ttk.Entry(info_frame, width=30)
        entries["Name"].grid(row=0, column=1, columnspan=3, sticky="w", **pad)

        ttk.Label(info_frame, text="Phone:").grid(row=1, column=0, sticky="e", **pad)
        entries["Phone"] = ttk.Entry(info_frame, width=18)
        entries["Phone"].grid(row=1, column=1, sticky="w", **pad)

        ttk.Label(info_frame, text="Email:").grid(row=1, column=2, sticky="e", **pad)
        entries["Email"] = ttk.Entry(info_frame, width=20)
        entries["Email"].grid(row=1, column=3, sticky="w", **pad)

        ttk.Label(info_frame, text="EIN:").grid(row=2, column=0, sticky="e", **pad)
        entries["EIN Number"] = ttk.Entry(info_frame, width=18)
        entries["EIN Number"].grid(row=2, column=1, sticky="w", **pad)

        # Address section
        addr_frame = ttk.LabelFrame(dialog, text="Address", padding=10)
        addr_frame.pack(fill="x", padx=15, pady=5)

        ttk.Label(addr_frame, text="Street:").grid(row=0, column=0, sticky="e", **pad)
        entries["Street Address"] = ttk.Entry(addr_frame, width=35)
        entries["Street Address"].grid(row=0, column=1, columnspan=3, sticky="w", **pad)

        ttk.Label(addr_frame, text="Zip:").grid(row=1, column=0, sticky="e", **pad)
        zip_entry = ttk.Entry(addr_frame, width=8)
        zip_entry.grid(row=1, column=1, sticky="w", **pad)
        entries["Zip Code"] = zip_entry

        ttk.Label(addr_frame, text="City:").grid(row=1, column=2, sticky="e", **pad)
        city_entry = ttk.Entry(addr_frame, width=15)
        city_entry.grid(row=1, column=3, sticky="w", **pad)
        entries["City"] = city_entry

        ttk.Label(addr_frame, text="State:").grid(row=2, column=0, sticky="e", **pad)
        state_entry = ttk.Entry(addr_frame, width=5)
        state_entry.grid(row=2, column=1, sticky="w", **pad)
        entries["State"] = state_entry

        zip_entry.bind("<FocusOut>", lambda e: self._lookup_zip(zip_entry, city_entry, state_entry))
        zip_entry.bind("<Return>", lambda e: self._lookup_zip(zip_entry, city_entry, state_entry))

        # Save button
        def save():
            required = {
                "Name": entries["Name"].get().strip(),
                "Phone": entries["Phone"].get().strip(),
                "Email": entries["Email"].get().strip(),
                "EIN Number": entries["EIN Number"].get().strip(),
                "Street Address": entries["Street Address"].get().strip(),
                "Zip Code": entries["Zip Code"].get().strip(),
                "City": entries["City"].get().strip(),
                "State": entries["State"].get().strip(),
            }
            missing = [k for k, v in required.items() if not v]
            if missing:
                messagebox.showerror("Error", f"All fields are required.\nMissing: {', '.join(missing)}")
                return
            code = db.add_contractor(
                "Primary",
                required["Name"],
                street=required["Street Address"],
                city=required["City"],
                state=required["State"],
                zipcode=required["Zip Code"],
                phone=required["Phone"],
                email=required["Email"],
                ein=required["EIN Number"],
                user_id=self.user_id,
            )
            dialog.destroy()
            self._refresh_primary()
            messagebox.showinfo("Added", f"Primary contractor added.\nCode ID: {code}")

        ttk.Button(dialog, text="Save", command=save).pack(pady=15)

    def _add_sub_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Add Subcontractor")
        dialog.geometry("480x420")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 12, "pady": 6}
        entries = {}

        # Info section
        info_frame = ttk.LabelFrame(dialog, text="Subcontractor Info", padding=10)
        info_frame.pack(fill="x", padx=15, pady=(15, 5))

        ttk.Label(info_frame, text="Name:").grid(row=0, column=0, sticky="e", **pad)
        entries["Name"] = ttk.Entry(info_frame, width=30)
        entries["Name"].grid(row=0, column=1, columnspan=3, sticky="w", **pad)

        ttk.Label(info_frame, text="Phone:").grid(row=1, column=0, sticky="e", **pad)
        entries["Phone"] = ttk.Entry(info_frame, width=18)
        entries["Phone"].grid(row=1, column=1, sticky="w", **pad)

        ttk.Label(info_frame, text="Email:").grid(row=1, column=2, sticky="e", **pad)
        entries["Email"] = ttk.Entry(info_frame, width=20)
        entries["Email"].grid(row=1, column=3, sticky="w", **pad)

        ttk.Label(info_frame, text="SSN/TIN:").grid(row=2, column=0, sticky="e", **pad)
        entries["SSN or TIN"] = ttk.Entry(info_frame, width=18)
        entries["SSN or TIN"].grid(row=2, column=1, sticky="w", **pad)

        # Address section
        addr_frame = ttk.LabelFrame(dialog, text="Address", padding=10)
        addr_frame.pack(fill="x", padx=15, pady=5)

        ttk.Label(addr_frame, text="Street:").grid(row=0, column=0, sticky="e", **pad)
        entries["Street Address"] = ttk.Entry(addr_frame, width=35)
        entries["Street Address"].grid(row=0, column=1, columnspan=3, sticky="w", **pad)

        ttk.Label(addr_frame, text="Zip:").grid(row=1, column=0, sticky="e", **pad)
        zip_entry = ttk.Entry(addr_frame, width=8)
        zip_entry.grid(row=1, column=1, sticky="w", **pad)
        entries["Zip Code"] = zip_entry

        ttk.Label(addr_frame, text="City:").grid(row=1, column=2, sticky="e", **pad)
        city_entry = ttk.Entry(addr_frame, width=15)
        city_entry.grid(row=1, column=3, sticky="w", **pad)
        entries["City"] = city_entry

        ttk.Label(addr_frame, text="State:").grid(row=2, column=0, sticky="e", **pad)
        state_entry = ttk.Entry(addr_frame, width=5)
        state_entry.grid(row=2, column=1, sticky="w", **pad)
        entries["State"] = state_entry

        zip_entry.bind("<FocusOut>", lambda e: self._lookup_zip(zip_entry, city_entry, state_entry))
        zip_entry.bind("<Return>", lambda e: self._lookup_zip(zip_entry, city_entry, state_entry))

        # Save button
        def save():
            required = {
                "Name": entries["Name"].get().strip(),
                "Phone": entries["Phone"].get().strip(),
                "Email": entries["Email"].get().strip(),
                "SSN or TIN": entries["SSN or TIN"].get().strip(),
                "Street Address": entries["Street Address"].get().strip(),
                "Zip Code": entries["Zip Code"].get().strip(),
                "City": entries["City"].get().strip(),
                "State": entries["State"].get().strip(),
            }
            missing = [k for k, v in required.items() if not v]
            if missing:
                messagebox.showerror("Error", f"All fields are required.\nMissing: {', '.join(missing)}")
                return
            code = db.add_contractor(
                "Subcontractor",
                required["Name"],
                street=required["Street Address"],
                city=required["City"],
                state=required["State"],
                zipcode=required["Zip Code"],
                phone=required["Phone"],
                email=required["Email"],
                ssn_tin=required["SSN or TIN"],
                user_id=self.user_id,
            )
            dialog.destroy()
            self._refresh_subs()
            messagebox.showinfo("Added", f"Subcontractor added.\nCode ID: {code}")

        ttk.Button(dialog, text="Save", command=save).pack(pady=15)

    def _delete_primary(self):
        selected = self.primary_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select a primary contractor to delete.")
            return
        name = self.primary_tree.item(selected, "values")[1]
        if messagebox.askyesno("Confirm", f"Delete primary contractor '{name}'?"):
            db.delete_contractor(int(selected))
            self._refresh_primary()

    def _delete_sub(self):
        selected = self.sub_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select a subcontractor to delete.")
            return
        name = self.sub_tree.item(selected, "values")[1]
        if messagebox.askyesno("Confirm", f"Delete subcontractor '{name}'?"):
            db.delete_contractor(int(selected))
            self._refresh_subs()

    # ──────────────────────────────────────────────
    # Invoicing Tab
    # ──────────────────────────────────────────────
    def _build_invoicing_tab(self):
        frame = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(frame, text="Invoicing")

        self.invoice_notebook = ttk.Notebook(frame)
        self.invoice_notebook.pack(fill="both", expand=True)

        self._build_primary_invoices_tab()
        self._build_client_invoices_tab()
        self._build_clients_tab()
        self._build_services_tab()

    def _build_primary_invoices_tab(self):
        frame = ttk.Frame(self.invoice_notebook, padding=10)
        self.invoice_notebook.add(frame, text="Primary Invoices")

        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))
        ttk.Button(toolbar, text="New Weekly Bill", command=self._new_primary_invoice).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Open", command=self._view_primary_invoice).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete", command=lambda: self._delete_invoice("Primary")).pack(side="right", padx=5)

        # Filter by Primary
        filter_frame = ttk.Frame(frame)
        filter_frame.pack(fill="x", pady=(0, 5))
        ttk.Label(filter_frame, text="Primary:").pack(side="left", padx=(0, 5))
        self._primary_inv_filter_var = tk.StringVar(value="All")
        self._primary_inv_combo = ttk.Combobox(filter_frame, textvariable=self._primary_inv_filter_var, width=25, state="readonly")
        self._primary_inv_combo.pack(side="left", padx=5)
        self._primary_inv_combo.bind("<<ComboboxSelected>>", lambda e: self._refresh_primary_invoices())
        self._update_primary_filter_list()

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("Invoice #", "Primary", "Week #", "From", "To", "Total", "Status", "Payment")
        self.primary_inv_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=14)
        for col in cols:
            self.primary_inv_tree.heading(col, text=col, anchor="center")
        self.primary_inv_tree.column("Invoice #", width=75, anchor="center", stretch=True)
        self.primary_inv_tree.column("Primary", width=155, anchor="center", stretch=True)
        self.primary_inv_tree.column("Week #", width=55, anchor="center", stretch=True)
        self.primary_inv_tree.column("From", width=85, anchor="center", stretch=True)
        self.primary_inv_tree.column("To", width=85, anchor="center", stretch=True)
        self.primary_inv_tree.column("Total", width=90, anchor="center", stretch=True)
        self.primary_inv_tree.column("Status", width=65, anchor="center", stretch=True)
        self.primary_inv_tree.column("Payment", width=100, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.primary_inv_tree.yview)
        self.primary_inv_tree.configure(yscrollcommand=scrollbar.set)
        self.primary_inv_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self._setup_sortable(self.primary_inv_tree, self._refresh_primary_invoices)
        self.primary_inv_tree.bind("<Double-1>", lambda e: self._view_primary_invoice())
        self._refresh_primary_invoices()

    def _update_primary_filter_list(self):
        primaries = db.get_contractors("Primary", user_id=self.user_id)
        names = ["All"] + [c[2] for c in primaries]
        self._primary_inv_combo["values"] = names
        self._primary_inv_filter_ids = {c[2]: c[0] for c in primaries}
        if not hasattr(self, "_primary_inv_filter_ids"):
            self._primary_inv_filter_ids = {}



    def _build_client_invoices_tab(self):
        frame = ttk.Frame(self.invoice_notebook, padding=10)
        self.invoice_notebook.add(frame, text="Client Invoices")

        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))
        ttk.Button(toolbar, text="New Invoice", command=self._new_client_invoice).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Open", command=self._view_client_invoice).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete", command=lambda: self._delete_invoice("Client")).pack(side="right", padx=5)

        # Filter by Client
        filter_frame = ttk.Frame(frame)
        filter_frame.pack(fill="x", pady=(0, 5))
        ttk.Label(filter_frame, text="Client:").pack(side="left", padx=(0, 5))
        self._client_inv_filter_var = tk.StringVar(value="All")
        self._client_inv_combo = ttk.Combobox(filter_frame, textvariable=self._client_inv_filter_var, width=25, state="readonly")
        self._client_inv_combo.pack(side="left", padx=5)
        self._client_inv_combo.bind("<<ComboboxSelected>>", lambda e: self._refresh_client_invoices())
        self._update_client_filter_list()

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("Invoice #", "Client", "From", "To", "Total", "Status", "Payment")
        self.client_inv_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=14)
        for col in cols:
            self.client_inv_tree.heading(col, text=col, anchor="center")
        self.client_inv_tree.column("Invoice #", width=80, anchor="center", stretch=True)
        self.client_inv_tree.column("Client", width=160, anchor="center", stretch=True)
        self.client_inv_tree.column("From", width=85, anchor="center", stretch=True)
        self.client_inv_tree.column("To", width=85, anchor="center", stretch=True)
        self.client_inv_tree.column("Total", width=90, anchor="center", stretch=True)
        self.client_inv_tree.column("Status", width=65, anchor="center", stretch=True)
        self.client_inv_tree.column("Payment", width=100, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.client_inv_tree.yview)
        self.client_inv_tree.configure(yscrollcommand=scrollbar.set)
        self.client_inv_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self._setup_sortable(self.client_inv_tree, self._refresh_client_invoices)
        self.client_inv_tree.bind("<Double-1>", lambda e: self._view_client_invoice())
        self._refresh_client_invoices()

    def _update_client_filter_list(self):
        clients = db.get_clients(user_id=self.user_id)
        names = ["All"] + [c[2] for c in clients]
        self._client_inv_combo["values"] = names
        self._client_inv_filter_ids = {c[2]: c[0] for c in clients}

    def _build_clients_tab(self):
        frame = ttk.Frame(self.invoice_notebook, padding=10)
        self.invoice_notebook.add(frame, text="Individual Clients")

        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))
        ttk.Button(toolbar, text="Add Client", command=self._add_client_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete Selected", command=self._delete_client).pack(side="left", padx=5)

        search_frame = ttk.Frame(frame)
        search_frame.pack(fill="x", pady=(0, 5))
        ttk.Label(search_frame, text="Search:").pack(side="left", padx=(0, 5))
        self._client_filter_var = tk.StringVar()
        ttk.Entry(search_frame, textvariable=self._client_filter_var, width=30).pack(side="left", padx=(0, 5))
        ttk.Button(search_frame, text="Clear", command=lambda: self._client_filter_var.set("")).pack(side="left")
        self._client_filter_var.trace_add("write", lambda *a: self._refresh_clients())

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("ID", "Name", "Address", "City", "State", "Zip", "Phone")
        self.client_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=14)
        for col in cols:
            self.client_tree.heading(col, text=col, anchor="center")
        self.client_tree.column("ID", width=50, anchor="center", stretch=True)
        self.client_tree.column("Name", width=180, anchor="center", stretch=True)
        self.client_tree.column("Address", width=160, anchor="center", stretch=True)
        self.client_tree.column("City", width=100, anchor="center", stretch=True)
        self.client_tree.column("State", width=45, anchor="center", stretch=True)
        self.client_tree.column("Zip", width=55, anchor="center", stretch=True)
        self.client_tree.column("Phone", width=110, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.client_tree.yview)
        self.client_tree.configure(yscrollcommand=scrollbar.set)
        self.client_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self._setup_sortable(self.client_tree, self._refresh_clients)
        self._refresh_clients()

    # ── Invoicing Refresh ──

    def _refresh_primary_invoices(self):
        self._update_primary_filter_list()
        for row in self.primary_inv_tree.get_children():
            self.primary_inv_tree.delete(row)
        filter_name = self._primary_inv_filter_var.get() if hasattr(self, "_primary_inv_filter_var") else "All"
        filter_id = self._primary_inv_filter_ids.get(filter_name) if filter_name != "All" else None
        invoices = db.get_invoices("Primary", user_id=self.user_id)
        for inv in invoices:
            iid, inv_num, inv_type, rec_type, rec_id, week, dfrom, dto, created, status, total = inv[:11]
            if filter_id and rec_id != filter_id:
                continue
            name = db.get_recipient_name(rec_type, rec_id)
            # Get payment info
            full_inv = db.get_invoice(iid)
            payment_method = full_inv[10] or "" if full_inv else ""
            full_inv_notes = full_inv[11] or "" if full_inv else ""
            payment_display = full_inv_notes if status == "Paid" else ""
            values = (inv_num, name, week or "", dfrom or "", dto or "", f"${total:,.2f}", status, payment_display)
            self.primary_inv_tree.insert("", "end", iid=iid, values=values)

    def _refresh_client_invoices(self):
        self._update_client_filter_list()
        for row in self.client_inv_tree.get_children():
            self.client_inv_tree.delete(row)
        filter_name = self._client_inv_filter_var.get() if hasattr(self, "_client_inv_filter_var") else "All"
        filter_id = self._client_inv_filter_ids.get(filter_name) if filter_name != "All" else None
        invoices = db.get_invoices("Client", user_id=self.user_id)
        for inv in invoices:
            iid, inv_num, inv_type, rec_type, rec_id, week, dfrom, dto, created, status, total = inv[:11]
            if filter_id and rec_id != filter_id:
                continue
            name = db.get_recipient_name(rec_type, rec_id)
            full_inv = db.get_invoice(iid)
            payment_notes = full_inv[11] or "" if full_inv else ""
            payment_display = payment_notes if status == "Paid" else ""
            values = (inv_num, name, dfrom or "", dto or "", f"${total:,.2f}", status, payment_display)
            self.client_inv_tree.insert("", "end", iid=iid, values=values)

    def _refresh_clients(self):
        for row in self.client_tree.get_children():
            self.client_tree.delete(row)
        query = self._client_filter_var.get().lower() if hasattr(self, "_client_filter_var") else ""
        cols = self.client_tree["columns"]
        clients = db.get_clients(user_id=self.user_id)
        for c in clients:
            cid, code_id, name, street, city, state, zipcode, phone = c
            values = (code_id, name, street or "", city or "", state or "", zipcode or "", phone or "")
            if query and not any(query in str(v).lower() for v in values):
                continue
            if not self._passes_filters(self.client_tree, values, cols):
                continue
            self.client_tree.insert("", "end", iid=cid, values=values)

    # ── Client Management ──

    def _add_client_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Add Individual Client")
        dialog.geometry("450x400")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 12, "pady": 6}
        entries = {}

        info_frame = ttk.LabelFrame(dialog, text="Client Info", padding=10)
        info_frame.pack(fill="x", padx=15, pady=(15, 5))

        ttk.Label(info_frame, text="Name:").grid(row=0, column=0, sticky="e", **pad)
        entries["Name"] = ttk.Entry(info_frame, width=30)
        entries["Name"].grid(row=0, column=1, columnspan=3, sticky="w", **pad)

        ttk.Label(info_frame, text="Phone:").grid(row=1, column=0, sticky="e", **pad)
        entries["Phone"] = ttk.Entry(info_frame, width=18)
        entries["Phone"].grid(row=1, column=1, sticky="w", **pad)

        addr_frame = ttk.LabelFrame(dialog, text="Address", padding=10)
        addr_frame.pack(fill="x", padx=15, pady=5)

        ttk.Label(addr_frame, text="Street:").grid(row=0, column=0, sticky="e", **pad)
        entries["Street"] = ttk.Entry(addr_frame, width=35)
        entries["Street"].grid(row=0, column=1, columnspan=3, sticky="w", **pad)

        ttk.Label(addr_frame, text="Zip:").grid(row=1, column=0, sticky="e", **pad)
        zip_entry = ttk.Entry(addr_frame, width=8)
        zip_entry.grid(row=1, column=1, sticky="w", **pad)
        entries["Zip"] = zip_entry

        ttk.Label(addr_frame, text="City:").grid(row=1, column=2, sticky="e", **pad)
        city_entry = ttk.Entry(addr_frame, width=15)
        city_entry.grid(row=1, column=3, sticky="w", **pad)
        entries["City"] = city_entry

        ttk.Label(addr_frame, text="State:").grid(row=2, column=0, sticky="e", **pad)
        state_entry = ttk.Entry(addr_frame, width=5)
        state_entry.grid(row=2, column=1, sticky="w", **pad)
        entries["State"] = state_entry

        zip_entry.bind("<FocusOut>", lambda e: self._lookup_zip(zip_entry, city_entry, state_entry))
        zip_entry.bind("<Return>", lambda e: self._lookup_zip(zip_entry, city_entry, state_entry))

        def save():
            name = entries["Name"].get().strip()
            if not name:
                messagebox.showerror("Error", "Name is required.")
                return
            code = db.add_client(
                name,
                street=entries["Street"].get().strip(),
                city=entries["City"].get().strip(),
                state=entries["State"].get().strip(),
                zipcode=entries["Zip"].get().strip(),
                phone=entries["Phone"].get().strip(),
                user_id=self.user_id,
            )
            dialog.destroy()
            self._refresh_clients()
            messagebox.showinfo("Added", f"Client added.\nCode ID: {code}")

        ttk.Button(dialog, text="Save", command=save).pack(pady=15)

    def _delete_client(self):
        selected = self.client_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select a client to delete.")
            return
        name = self.client_tree.item(selected, "values")[1]
        if messagebox.askyesno("Confirm", f"Delete client '{name}'?"):
            db.delete_client(int(selected))
            self._refresh_clients()

    # ── Invoice Creation ──

    def _new_primary_invoice(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("New Weekly Bill")
        dialog.geometry("500x250")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 12, "pady": 8}

        ttk.Label(dialog, text="Create Weekly Bill for Primary Contractor", font=("Segoe UI", 11, "bold")).pack(pady=(15, 10))

        form = ttk.Frame(dialog)
        form.pack(fill="x", padx=20)

        ttk.Label(form, text="Primary:").grid(row=0, column=0, sticky="e", **pad)
        primaries = db.get_contractors("Primary", user_id=self.user_id)
        primary_names = [f"{c[1]} - {c[2]}" for c in primaries]
        primary_ids = [c[0] for c in primaries]
        primary_var = tk.StringVar()
        primary_combo = ttk.Combobox(form, textvariable=primary_var, values=primary_names, width=30, state="readonly")
        primary_combo.grid(row=0, column=1, sticky="w", **pad)

        ttk.Label(form, text="Week #:").grid(row=1, column=0, sticky="e", **pad)
        week_entry = ttk.Entry(form, width=8)
        week_entry.grid(row=1, column=1, sticky="w", **pad)

        ttk.Label(form, text="From:").grid(row=2, column=0, sticky="e", **pad)
        from_frame, from_var = self._create_date_entry(form)
        from_frame.grid(row=2, column=1, sticky="w", **pad)

        ttk.Label(form, text="To:").grid(row=2, column=2, sticky="e", **pad)
        to_frame, to_var = self._create_date_entry(form)
        to_frame.grid(row=2, column=3, sticky="w", **pad)

        def save():
            idx = primary_combo.current()
            if idx < 0:
                messagebox.showerror("Error", "Select a primary contractor.")
                return
            week = week_entry.get().strip()
            inv_id, inv_num = db.create_invoice(
                "Primary", "Primary", primary_ids[idx],
                week_number=int(week) if week.isdigit() else None,
                date_from=from_var.get(),
                date_to=to_var.get(),
                user_id=self.user_id,
            )
            dialog.destroy()
            self._refresh_primary_invoices()
            self._open_invoice_editor(inv_id, inv_num)

        ttk.Button(dialog, text="Create & Open", command=save).pack(pady=15)

    def _new_client_invoice(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("New Client Invoice")
        dialog.geometry("500x220")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 12, "pady": 8}

        ttk.Label(dialog, text="Create Invoice for Individual Client", font=("Segoe UI", 11, "bold")).pack(pady=(15, 10))

        form = ttk.Frame(dialog)
        form.pack(fill="x", padx=20)

        ttk.Label(form, text="Client:").grid(row=0, column=0, sticky="e", **pad)
        clients = db.get_clients(user_id=self.user_id)
        client_names = [f"{c[1]} - {c[2]}" for c in clients]
        client_ids = [c[0] for c in clients]
        client_var = tk.StringVar()
        client_combo = ttk.Combobox(form, textvariable=client_var, values=client_names, width=30, state="readonly")
        client_combo.grid(row=0, column=1, columnspan=3, sticky="w", **pad)

        ttk.Label(form, text="From:").grid(row=1, column=0, sticky="e", **pad)
        from_frame, from_var = self._create_date_entry(form)
        from_frame.grid(row=1, column=1, sticky="w", **pad)

        ttk.Label(form, text="To:").grid(row=1, column=2, sticky="e", **pad)
        to_frame, to_var = self._create_date_entry(form)
        to_frame.grid(row=1, column=3, sticky="w", **pad)

        def save():
            idx = client_combo.current()
            if idx < 0:
                messagebox.showerror("Error", "Select a client.")
                return
            inv_id, inv_num = db.create_invoice(
                "Client", "Client", client_ids[idx],
                date_from=from_var.get(),
                date_to=to_var.get(),
                user_id=self.user_id,
            )
            dialog.destroy()
            self._refresh_client_invoices()
            self._open_invoice_editor(inv_id, inv_num)

        ttk.Button(dialog, text="Create & Open", command=save).pack(pady=15)

    # ── Invoice Workspace ──

    def _open_invoice_editor(self, invoice_id, inv_num):
        dialog = tk.Toplevel(self.root)
        dialog.geometry("850x680")
        dialog.transient(self.root)

        inv = db.get_invoice(invoice_id)
        _, _, inv_type, rec_type, rec_id, week, dfrom, dto, created, status, pay_method, pay_notes, total_val = inv
        recipient_name = db.get_recipient_name(rec_type, rec_id)
        is_primary = (inv_type == "Primary")

        if is_primary:
            dialog.title(f"Weekly Invoice: {inv_num}")
            header_text = f"{inv_num}  •  {recipient_name}  •  Week {week or ''}"
        else:
            dialog.title(f"Client Invoice: {inv_num}")
            header_text = f"{inv_num}  •  {recipient_name}"

        # Header
        header = ttk.Frame(dialog)
        header.pack(fill="x", padx=15, pady=(12, 5))
        ttk.Label(header, text=header_text, font=("Segoe UI", 13, "bold")).pack(side="left")
        ttk.Label(header, text=f"{dfrom or ''} → {dto or ''}", font=("Segoe UI", 10)).pack(side="left", padx=15)

        # Payment row
        pay_frame = ttk.Frame(dialog)
        pay_frame.pack(fill="x", padx=15, pady=(0, 5))

        paid_var = tk.BooleanVar(value=(status == "Paid"))

        def save_payment(*args):
            if paid_var.get():
                db.update_invoice_status(invoice_id, "Paid", None, notes_entry.get().strip())
            else:
                db.update_invoice_status(invoice_id, "Unpaid", None, None)
            if is_primary:
                self._refresh_primary_invoices()
            else:
                self._refresh_client_invoices()

        ttk.Checkbutton(pay_frame, text="Paid", variable=paid_var, command=save_payment).pack(side="left", padx=3)
        ttk.Label(pay_frame, text="Notes:").pack(side="left", padx=(10, 2))
        notes_entry = ttk.Entry(pay_frame, width=30)
        notes_entry.insert(0, pay_notes or "")
        notes_entry.pack(side="left", padx=2)
        notes_entry.bind("<FocusOut>", save_payment)

        # All jobs tree
        tree_frame = ttk.Frame(dialog)
        tree_frame.pack(fill="both", expand=True, padx=15, pady=8)

        cols = ("Date", "Customer", "Mobile #", "Service", "Price")
        job_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=18)
        job_tree.heading("Date", text="Date", anchor="center")
        job_tree.heading("Customer", text="Customer", anchor="center")
        job_tree.heading("Mobile #", text="Mobile #", anchor="center")
        job_tree.heading("Service", text="Service", anchor="center")
        job_tree.heading("Price", text="Price", anchor="center")
        job_tree.column("Date", width=90, anchor="center", stretch=True)
        job_tree.column("Customer", width=160, anchor="center", stretch=True)
        job_tree.column("Mobile #", width=120, anchor="center", stretch=True)
        job_tree.column("Service", width=220, anchor="center", stretch=True)
        job_tree.column("Price", width=100, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=job_tree.yview)
        job_tree.configure(yscrollcommand=scrollbar.set)
        job_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Total
        total_frame = ttk.Frame(dialog)
        total_frame.pack(fill="x", padx=15)
        total_label = "Week Total" if is_primary else "Total"
        week_total_var = tk.StringVar(value=f"{total_label}: $0.00")
        ttk.Label(total_frame, textvariable=week_total_var, font=("Segoe UI", 12, "bold")).pack(side="right")

        def refresh_all():
            for row in job_tree.get_children():
                job_tree.delete(row)
            jobs = db.get_invoice_jobs(invoice_id)
            grand = 0
            for job in jobs:
                jid, jdate, cust, mobile, jtotal = job
                lines = db.get_job_lines(jid)
                if lines:
                    for i, ln in enumerate(lines):
                        lid, svc, price = ln
                        if i == 0:
                            job_tree.insert("", "end", tags=("job",), values=(jdate or "", cust or "", mobile or "", svc, f"${price:,.2f}"))
                        else:
                            job_tree.insert("", "end", values=("", "", "", svc, f"${price:,.2f}"))
                    job_tree.insert("", "end", tags=("total",), values=("", "", "", "Job Total", f"${jtotal:,.2f}"))
                    job_tree.insert("", "end", values=("", "", "", "", ""))
                else:
                    job_tree.insert("", "end", tags=("job",), values=(jdate or "", cust or "", mobile or "", "(no services)", "$0.00"))
                grand += jtotal
            week_total_var.set(f"{total_label}: ${grand:,.2f}")
            job_tree.tag_configure("total", font=("Segoe UI", 9, "bold"))

        def add_customer():
            cust_dialog = tk.Toplevel(dialog)
            cust_dialog.title("Add Customer Job")
            cust_dialog.geometry("600x560")
            cust_dialog.transient(dialog)
            cust_dialog.grab_set()
            self._center_dialog(cust_dialog)
            cust_dialog.resizable(False, False)

            # Customer info
            info = ttk.LabelFrame(cust_dialog, text="Customer", padding=10)
            info.pack(fill="x", padx=12, pady=(12, 5))

            ttk.Label(info, text="Name:").grid(row=0, column=0, sticky="e", padx=6, pady=5)
            name_e = ttk.Entry(info, width=22)
            name_e.grid(row=0, column=1, sticky="w", padx=6, pady=5)

            ttk.Label(info, text="Mobile #:").grid(row=0, column=2, sticky="e", padx=6, pady=5)
            mobile_e = ttk.Entry(info, width=16)
            mobile_e.grid(row=0, column=3, sticky="w", padx=6, pady=5)
            self._format_phone_entry(mobile_e)

            ttk.Label(info, text="Date:").grid(row=1, column=0, sticky="e", padx=6, pady=5)
            min_d = datetime.strptime(dfrom, "%Y-%m-%d").date() if dfrom else None
            max_d = datetime.strptime(dto, "%Y-%m-%d").date() if dto else None
            date_frame, date_var = self._create_date_entry(info, default=dfrom, min_date=min_d, max_date=max_d)
            date_frame.grid(row=1, column=1, sticky="w", padx=6, pady=5)

            # Inline service add
            add_frame = ttk.LabelFrame(cust_dialog, text="Add Service", padding=8)
            add_frame.pack(fill="x", padx=12, pady=5)

            all_svcs = db.get_services(user_id=self.user_id)
            svc_names = [s[1] for s in all_svcs]
            svc_var = tk.StringVar()
            ttk.Label(add_frame, text="Service:").pack(side="left", padx=(0, 5))
            svc_combo = ttk.Combobox(add_frame, textvariable=svc_var, values=svc_names, width=22, state="readonly")
            svc_combo.pack(side="left", padx=5)

            ttk.Label(add_frame, text="$").pack(side="left", padx=(10, 2))
            price_e = ttk.Entry(add_frame, width=10)
            price_e.pack(side="left", padx=2)

            temp_services = []

            def add_svc(*args):
                s = svc_var.get().strip()
                p_str = price_e.get().strip().replace("$", "").replace(",", "")
                if not s:
                    return
                try:
                    p = float(p_str) if p_str else 0
                except ValueError:
                    return
                temp_services.append((s, p))
                svc_list.insert("", "end", values=(s, f"${p:,.2f}"))
                svc_var.set("")
                price_e.delete(0, "end")
                update_total()
                svc_combo.focus_set()

            ttk.Button(add_frame, text="Add", command=add_svc, width=5).pack(side="left", padx=8)
            price_e.bind("<Return>", add_svc)

            # Services list
            svc_frame = ttk.LabelFrame(cust_dialog, text="Services Added", padding=8)
            svc_frame.pack(fill="both", expand=True, padx=12, pady=5)

            svc_list = ttk.Treeview(svc_frame, columns=("Service", "Price"), show="headings", height=7)
            svc_list.heading("Service", text="Service", anchor="center")
            svc_list.heading("Price", text="Price", anchor="center")
            svc_list.column("Service", width=320, anchor="center", stretch=True)
            svc_list.column("Price", width=100, anchor="center", stretch=True)
            svc_list.pack(fill="both", expand=True)

            total_row = ttk.Frame(svc_frame)
            total_row.pack(fill="x", pady=(5, 0))
            svc_total_var = tk.StringVar(value="Total: $0.00")
            ttk.Label(total_row, textvariable=svc_total_var, font=("Segoe UI", 10, "bold")).pack(side="right")

            def update_total():
                t = sum(p for _, p in temp_services)
                svc_total_var.set(f"Total: ${t:,.2f}")

            def remove_svc():
                sel = svc_list.focus()
                if sel:
                    idx = svc_list.index(sel)
                    svc_list.delete(sel)
                    temp_services.pop(idx)
                    update_total()

            def save_customer():
                name = name_e.get().strip()
                if not name:
                    messagebox.showerror("Error", "Enter customer name.", parent=cust_dialog)
                    return
                if not temp_services:
                    messagebox.showerror("Error", "Add at least one service.", parent=cust_dialog)
                    return
                jid = db.add_invoice_job(invoice_id, date=date_var.get().strip(), customer_name=name, mobile_number=mobile_e.get().strip())
                for svc, price in temp_services:
                    db.add_invoice_line(jid, svc, price)
                db.update_job_total(jid)
                db.update_invoice_total(invoice_id)
                cust_dialog.destroy()
                refresh_all()
                if is_primary:
                    self._refresh_primary_invoices()
                else:
                    self._refresh_client_invoices()

            btn = ttk.Frame(cust_dialog)
            btn.pack(fill="x", padx=12, pady=8)
            ttk.Button(btn, text="Remove Selected", command=remove_svc).pack(side="left", padx=4)
            ttk.Button(btn, text="Save Customer", command=save_customer).pack(side="right", padx=4)

        def preview():
            db.update_invoice_total(invoice_id)
            self._preview_weekly_invoice(invoice_id)

        # Bottom buttons
        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(fill="x", padx=15, pady=10)
        ttk.Button(btn_frame, text="Add Customer", command=add_customer).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Preview / Export", command=preview).pack(side="left", padx=5)
        def close_editor():
            db.update_invoice_total(invoice_id)
            if is_primary:
                self._refresh_primary_invoices()
            else:
                self._refresh_client_invoices()
            dialog.destroy()

        ttk.Button(btn_frame, text="Close", command=close_editor).pack(side="right", padx=5)

        refresh_all()

    # ── Invoice Preview & Export ──

    def _preview_weekly_invoice(self, invoice_id):
        inv = db.get_invoice(invoice_id)
        if not inv:
            return
        _, inv_num, inv_type, rec_type, rec_id, week, dfrom, dto, created, status, pay_method, pay_notes, total = inv
        recipient_name = db.get_recipient_name(rec_type, rec_id)
        is_primary = (inv_type == "Primary")
        company = db.load_company_info(self.user_id)
        jobs = db.get_invoice_jobs(invoice_id)

        preview = tk.Toplevel(self.root)
        preview.title(f"Invoice Preview: {inv_num}")
        preview.geometry("750x650")
        self._center_dialog(preview)

        # Canvas-based clean preview
        canvas = tk.Canvas(preview, bg="white")
        canvas.pack(fill="both", expand=True, padx=10, pady=(10, 5))

        y = 20
        # Company logo (top-left corner)
        logo_path = os.path.join(APP_DIR, "logo.png")
        if os.path.exists(logo_path):
            logo_img = PILImage.open(logo_path)
            logo_img.thumbnail((70, 70))
            self._preview_logo = ImageTk.PhotoImage(logo_img)
            canvas.create_image(60, y, image=self._preview_logo, anchor="nw")

        # Company header (to the right of logo)
        header_x = 145 if os.path.exists(logo_path) else 60
        company_name = company.get("Company Name", "").upper()
        canvas.create_text(header_x, y + 2, text=company_name, font=("Segoe UI", 18, "bold"), anchor="nw")
        y += 28
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company['Email'])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")
        if info_parts:
            canvas.create_text(header_x, y + 2, text="  |  ".join(info_parts), font=("Segoe UI", 10, "bold"), anchor="nw")
            y += 22

        y = max(y + 15, 105)
        canvas.create_line(50, y, 700, y, width=2)
        y += 15

        if is_primary:
            canvas.create_text(60, y, text=f"WORK FOR: {recipient_name}", font=("Segoe UI", 12, "bold"), anchor="w")
            y += 22
            canvas.create_text(60, y, text=f"Week #{week or ''}     From: {dfrom or ''}     To: {dto or ''}", font=("Segoe UI", 10), anchor="w")
        else:
            canvas.create_text(60, y, text=f"INVOICE TO: {recipient_name}", font=("Segoe UI", 12, "bold"), anchor="w")
            y += 22
            canvas.create_text(60, y, text=f"From: {dfrom or ''}     To: {dto or ''}", font=("Segoe UI", 10), anchor="w")
        y += 25
        canvas.create_line(50, y, 700, y, width=1)
        y += 12

        # Column headers
        canvas.create_text(60, y, text="DATE", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(145, y, text="CUSTOMER", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(310, y, text="MOBILE #", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(430, y, text="SERVICE", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(690, y, text="PRICE", font=("Segoe UI", 9, "bold"), anchor="e")
        y += 18
        canvas.create_line(50, y, 700, y, width=0.5)
        y += 8

        grand_total = 0
        for job in jobs:
            jid, jdate, cust, mobile, jtotal = job
            lines = db.get_job_lines(jid)
            for i, ln in enumerate(lines):
                lid, svc, price = ln
                if i == 0:
                    canvas.create_text(60, y, text=jdate or "", font=("Segoe UI", 9), anchor="w")
                    canvas.create_text(145, y, text=cust or "", font=("Segoe UI", 9), anchor="w")
                    canvas.create_text(310, y, text=mobile or "", font=("Segoe UI", 9), anchor="w")
                canvas.create_text(430, y, text=svc, font=("Segoe UI", 9), anchor="w")
                canvas.create_text(690, y, text=f"${price:,.2f}", font=("Segoe UI", 9), anchor="e")
                y += 16
            canvas.create_text(600, y, text="Job Total:", font=("Segoe UI", 9, "bold"), anchor="e")
            canvas.create_text(690, y, text=f"${jtotal:,.2f}", font=("Segoe UI", 9, "bold"), anchor="e")
            y += 20
            grand_total += jtotal

        y += 5
        canvas.create_line(50, y, 700, y, width=1.5)
        y += 12
        total_text = "WEEK TOTAL:" if is_primary else "TOTAL:"
        canvas.create_text(600, y, text=total_text, font=("Segoe UI", 11, "bold"), anchor="e")
        canvas.create_text(690, y, text=f"${grand_total:,.2f}", font=("Segoe UI", 11, "bold"), anchor="e")
        y += 25

        if status == "Paid":
            pay_text = "PAID"
            if pay_notes:
                pay_text += f" — {pay_notes}"
            canvas.create_text(60, y, text=pay_text, font=("Segoe UI", 9, "italic"), anchor="w", fill="green")

        canvas.configure(scrollregion=(0, 0, 750, max(y + 30, 600)))

        # Export buttons
        btn_frame = ttk.Frame(preview)
        btn_frame.pack(fill="x", padx=10, pady=8)

        def export_pdf():
            filepath = filedialog.asksaveasfilename(parent=preview, title="Save as PDF", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
            if filepath:
                self._export_weekly_invoice_pdf(invoice_id, filepath)
                messagebox.showinfo("Saved", f"PDF saved to:\n{filepath}", parent=preview)

        ttk.Button(btn_frame, text="Export PDF", command=export_pdf).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Print", command=lambda: self._print_invoice_from_id(invoice_id)).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Close", command=preview.destroy).pack(side="right", padx=5)

    def _print_invoice_from_id(self, invoice_id):
        import tempfile
        filepath = tempfile.mktemp(suffix=".pdf")
        self._export_weekly_invoice_pdf(invoice_id, filepath)
        os.startfile(filepath, "print")

    def _export_weekly_invoice_pdf(self, invoice_id, filepath):
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        inv = db.get_invoice(invoice_id)
        _, inv_num, inv_type, rec_type, rec_id, week, dfrom, dto, created, status, pay_method, pay_notes, total = inv
        recipient_name = db.get_recipient_name(rec_type, rec_id)
        is_primary = (inv_type == "Primary")
        company = db.load_company_info(self.user_id)
        jobs = db.get_invoice_jobs(invoice_id)

        doc = SimpleDocTemplate(filepath, pagesize=letter, topMargin=0.5 * inch)
        styles = getSampleStyleSheet()
        elements = []

        logo_path = os.path.join(APP_DIR, "logo.png")
        header_style = ParagraphStyle("H", parent=styles["Heading1"], fontSize=16, spaceAfter=4, fontName="Helvetica-Bold")
        sub_style = ParagraphStyle("S", parent=styles["Normal"], fontSize=10, spaceAfter=2, fontName="Helvetica-Bold")
        title_style = ParagraphStyle("T", parent=styles["Heading2"], fontSize=12, spaceAfter=6)

        company_name = company.get("Company Name", "").upper()
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company['Email'])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")

        if os.path.exists(logo_path):
            logo = Image(logo_path, width=0.9 * inch, height=0.9 * inch)
            header_content = []
            if company_name:
                header_content.append(Paragraph(company_name, header_style))
            if info_parts:
                header_content.append(Paragraph("  |  ".join(info_parts), sub_style))
            header_table_data = [[logo, header_content]]
            header_table = Table(header_table_data, colWidths=[1.1 * inch, 5.3 * inch])
            header_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))
            elements.append(header_table)
        else:
            if company_name:
                elements.append(Paragraph(company_name, header_style))
            if info_parts:
                elements.append(Paragraph("  |  ".join(info_parts), sub_style))
        elements.append(Spacer(1, 14))

        if is_primary:
            elements.append(Paragraph(f"WORK FOR: {recipient_name}", title_style))
            elements.append(Paragraph(f"Week #{week or ''}  |  From: {dfrom or ''}  |  To: {dto or ''}", sub_style))
        else:
            elements.append(Paragraph(f"INVOICE TO: {recipient_name}", title_style))
            elements.append(Paragraph(f"From: {dfrom or ''}  |  To: {dto or ''}", sub_style))
        elements.append(Spacer(1, 10))

        data = [["Date", "Customer", "Mobile #", "Service", "Price"]]
        grand_total = 0
        for job in jobs:
            jid, jdate, cust, mobile, jtotal = job
            job_lines = db.get_job_lines(jid)
            for i, ln in enumerate(job_lines):
                lid, svc, price = ln
                if i == 0:
                    data.append([jdate or "", cust or "", mobile or "", svc, f"${price:,.2f}"])
                else:
                    data.append(["", "", "", svc, f"${price:,.2f}"])
            data.append(["", "", "", "Job Total", f"${jtotal:,.2f}"])
            data.append(["", "", "", "", ""])
            grand_total += jtotal

        data.append(["", "", "", "WEEK TOTAL" if is_primary else "TOTAL", f"${grand_total:,.2f}"])

        table = Table(data, colWidths=[0.8*inch, 1.5*inch, 1.1*inch, 2*inch, 1*inch])
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (4, 0), (4, -1), "RIGHT"),
            ("LINEBELOW", (0, 0), (-1, 0), 1, colors.black),
            ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
            ("FONTNAME", (3, -1), (4, -1), "Helvetica-Bold"),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 15))

        if status == "Paid":
            pay_text = "PAID"
            if pay_notes:
                pay_text += f" — {pay_notes}"
            elements.append(Paragraph(pay_text, sub_style))

        doc.build(elements)

    # ── Services Tab ──

    def _build_services_tab(self):
        frame = ttk.Frame(self.invoice_notebook, padding=10)
        self.invoice_notebook.add(frame, text="Services")

        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))
        ttk.Button(toolbar, text="Add Service", command=self._add_service_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Edit Selected", command=self._edit_service_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete Selected", command=self._delete_service).pack(side="left", padx=5)

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("ID", "Service Name")
        self.service_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=16)
        self.service_tree.heading("ID", text="ID", anchor="center")
        self.service_tree.heading("Service Name", text="Service Name", anchor="center")
        self.service_tree.column("ID", width=50, anchor="center", stretch=True)
        self.service_tree.column("Service Name", width=400, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.service_tree.yview)
        self.service_tree.configure(yscrollcommand=scrollbar.set)
        self.service_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self._refresh_services()

    def _refresh_services(self):
        for row in self.service_tree.get_children():
            self.service_tree.delete(row)
        services = db.get_services(user_id=self.user_id)
        for s in services:
            self.service_tree.insert("", "end", iid=s[0], values=(s[0], s[1]))

    def _add_service_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Add Service")
        dialog.geometry("350x160")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        ttk.Label(dialog, text="Service Name:").pack(pady=(15, 5))
        name_entry = ttk.Entry(dialog, width=35)
        name_entry.pack(pady=5)

        def save():
            name = name_entry.get().strip()
            if not name:
                return
            if db.add_service(name, user_id=self.user_id):
                dialog.destroy()
                self._refresh_services()
            else:
                messagebox.showwarning("Exists", f"Service '{name}' already exists.")

        ttk.Button(dialog, text="Save", command=save).pack(pady=10)

    def _edit_service_dialog(self):
        selected = self.service_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select a service to edit.")
            return
        current_name = self.service_tree.item(selected, "values")[1]
        service_id = int(selected)

        dialog = tk.Toplevel(self.root)
        dialog.title("Edit Service")
        dialog.geometry("380x160")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        ttk.Label(dialog, text="Service Name:").pack(pady=(20, 5))
        name_entry = ttk.Entry(dialog, width=35)
        name_entry.insert(0, current_name)
        name_entry.pack(pady=5)
        name_entry.select_range(0, tk.END)
        name_entry.focus()

        def save():
            new_name = name_entry.get().strip()
            if not new_name:
                return
            if new_name == current_name:
                dialog.destroy()
                return
            db.update_service(service_id, new_name)
            dialog.destroy()
            self._refresh_services()

        ttk.Button(dialog, text="Save", command=save).pack(pady=10)

    def _delete_service(self):
        selected = self.service_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select a service to delete.")
            return
        name = self.service_tree.item(selected, "values")[1]
        if messagebox.askyesno("Confirm", f"Delete service '{name}'?"):
            db.delete_service(int(selected))
            self._refresh_services()

    # ── Invoice Actions ──

    def _view_primary_invoice(self):
        selected = self.primary_inv_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select an invoice to view.")
            return
        inv_num = self.primary_inv_tree.item(selected, "values")[0]
        self._open_invoice_editor(int(selected), inv_num)

    def _view_client_invoice(self):
        selected = self.client_inv_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select an invoice to view.")
            return
        inv_num = self.client_inv_tree.item(selected, "values")[0]
        self._open_invoice_editor(int(selected), inv_num)

    def _mark_invoice_status(self, inv_type, status):
        tree = self.primary_inv_tree if inv_type == "Primary" else self.client_inv_tree
        selected = tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select an invoice.")
            return
        db.update_invoice_status(int(selected), status)
        if inv_type == "Primary":
            self._refresh_primary_invoices()
        else:
            self._refresh_client_invoices()

    def _delete_invoice(self, inv_type):
        tree = self.primary_inv_tree if inv_type == "Primary" else self.client_inv_tree
        selected = tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select an invoice to delete.")
            return
        if messagebox.askyesno("Confirm", "Delete this invoice and all its jobs/lines?"):
            db.delete_invoice(int(selected))
            if inv_type == "Primary":
                self._refresh_primary_invoices()
            else:
                self._refresh_client_invoices()

    # ──────────────────────────────────────────────
    # Sub Payments Tab
    # ──────────────────────────────────────────────
    def _build_sub_payments_tab(self):
        frame = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(frame, text="Sub Payments")

        # Toolbar
        toolbar = ttk.Frame(frame)
        toolbar.pack(fill="x", pady=(0, 5))
        ttk.Button(toolbar, text="Record Payment", command=self._add_sub_payment_dialog).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Delete Selected", command=self._delete_sub_payment).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Individual Statement", command=self._generate_individual_sub_statement).pack(side="right", padx=5)
        ttk.Button(toolbar, text="Full Statement", command=self._generate_sub_statement).pack(side="right", padx=5)

        # Filter row
        filter_frame = ttk.Frame(frame)
        filter_frame.pack(fill="x", pady=(0, 5))

        ttk.Label(filter_frame, text="Sub:").pack(side="left", padx=(0, 5))
        self._sub_pay_filter_var = tk.StringVar(value="All")
        self._sub_pay_combo = ttk.Combobox(filter_frame, textvariable=self._sub_pay_filter_var, width=20, state="readonly")
        self._sub_pay_combo.pack(side="left", padx=5)
        self._sub_pay_combo.bind("<<ComboboxSelected>>", lambda e: self._refresh_sub_payments())

        ttk.Label(filter_frame, text="From:").pack(side="left", padx=(15, 5))
        self._sub_pay_from_var = tk.StringVar()
        ttk.Entry(filter_frame, textvariable=self._sub_pay_from_var, width=12).pack(side="left", padx=2)

        ttk.Label(filter_frame, text="To:").pack(side="left", padx=(10, 5))
        self._sub_pay_to_var = tk.StringVar()
        ttk.Entry(filter_frame, textvariable=self._sub_pay_to_var, width=12).pack(side="left", padx=2)

        ttk.Button(filter_frame, text="Filter", command=self._refresh_sub_payments).pack(side="left", padx=8)
        ttk.Button(filter_frame, text="Clear", command=self._clear_sub_pay_filter).pack(side="left", padx=2)

        # Treeview
        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill="both", expand=True)

        cols = ("Date", "Subcontractor", "Amount", "Type", "Period From", "Period To", "Notes")
        self.sub_pay_tree = ttk.Treeview(tree_frame, columns=cols, show="headings", height=16)
        self.sub_pay_tree.heading("Date", text="Date", anchor="center")
        self.sub_pay_tree.heading("Subcontractor", text="Subcontractor", anchor="center")
        self.sub_pay_tree.heading("Amount", text="Amount", anchor="center")
        self.sub_pay_tree.heading("Type", text="Type", anchor="center")
        self.sub_pay_tree.heading("Period From", text="Period From", anchor="center")
        self.sub_pay_tree.heading("Period To", text="Period To", anchor="center")
        self.sub_pay_tree.heading("Notes", text="Notes", anchor="center")
        self.sub_pay_tree.column("Date", width=90, anchor="center", stretch=True)
        self.sub_pay_tree.column("Subcontractor", width=150, anchor="center", stretch=True)
        self.sub_pay_tree.column("Amount", width=100, anchor="center", stretch=True)
        self.sub_pay_tree.column("Type", width=70, anchor="center", stretch=True)
        self.sub_pay_tree.column("Period From", width=90, anchor="center", stretch=True)
        self.sub_pay_tree.column("Period To", width=90, anchor="center", stretch=True)
        self.sub_pay_tree.column("Notes", width=150, anchor="center", stretch=True)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.sub_pay_tree.yview)
        self.sub_pay_tree.configure(yscrollcommand=scrollbar.set)
        self.sub_pay_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Total label
        self._sub_pay_total_var = tk.StringVar(value="Total: $0.00")
        ttk.Label(frame, textvariable=self._sub_pay_total_var, font=("Segoe UI", 11, "bold")).pack(anchor="e", padx=10, pady=(5, 0))

        self._setup_sortable(self.sub_pay_tree, self._refresh_sub_payments)
        self._update_sub_pay_filter_list()
        self._refresh_sub_payments()

    def _update_sub_pay_filter_list(self):
        subs = db.get_contractors("Subcontractor", user_id=self.user_id)
        names = ["All"] + [c[2] for c in subs]
        self._sub_pay_combo["values"] = names
        self._sub_pay_ids = {c[2]: c[0] for c in subs}

    def _clear_sub_pay_filter(self):
        self._sub_pay_filter_var.set("All")
        self._sub_pay_from_var.set("")
        self._sub_pay_to_var.set("")
        self._refresh_sub_payments()

    def _refresh_sub_payments(self):
        self._update_sub_pay_filter_list()
        for row in self.sub_pay_tree.get_children():
            self.sub_pay_tree.delete(row)

        filter_name = self._sub_pay_filter_var.get()
        contractor_id = self._sub_pay_ids.get(filter_name) if filter_name != "All" else None
        start = self._sub_pay_from_var.get().strip() or None
        end = self._sub_pay_to_var.get().strip() or None

        payments = db.get_sub_payments(start_date=start, end_date=end, contractor_id=contractor_id, user_id=self.user_id)
        total = 0
        for p in payments:
            pid, name, amount, pdate, ptype, pfrom, pto, notes = p
            values = (pdate, name, f"${amount:,.2f}", ptype, pfrom or "", pto or "", notes or "")
            self.sub_pay_tree.insert("", "end", iid=pid, values=values)
            total += amount
        self._sub_pay_total_var.set(f"Total: ${total:,.2f}")

    def _add_sub_payment_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Record Sub Payment")
        dialog.geometry("500x420")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        pad = {"padx": 10, "pady": 7}

        ttk.Label(dialog, text="Record Payment to Subcontractor", font=("Segoe UI", 11, "bold")).pack(pady=(15, 10))

        form = ttk.Frame(dialog)
        form.pack(fill="x", padx=20)

        # Subcontractor dropdown
        ttk.Label(form, text="Subcontractor:").grid(row=0, column=0, sticky="e", **pad)
        subs = db.get_contractors("Subcontractor", user_id=self.user_id)
        sub_names = [c[2] for c in subs]
        sub_ids = [c[0] for c in subs]
        sub_var = tk.StringVar()
        sub_combo = ttk.Combobox(form, textvariable=sub_var, values=sub_names, width=25, state="readonly")
        sub_combo.grid(row=0, column=1, columnspan=2, sticky="w", **pad)

        # Amount
        ttk.Label(form, text="Amount: $").grid(row=1, column=0, sticky="e", **pad)
        amt_entry = ttk.Entry(form, width=15)
        amt_entry.grid(row=1, column=1, sticky="w", **pad)

        # Payment Date
        ttk.Label(form, text="Payment Date:").grid(row=2, column=0, sticky="e", **pad)
        date_frame, date_var = self._create_date_entry(form)
        date_frame.grid(row=2, column=1, sticky="w", **pad)

        # Payment Type
        ttk.Label(form, text="Payment Type:").grid(row=3, column=0, sticky="e", **pad)
        type_var = tk.StringVar()
        type_combo = ttk.Combobox(form, textvariable=type_var, values=["Zelle", "Cash", "Check"], width=12, state="readonly")
        type_combo.grid(row=3, column=1, sticky="w", **pad)
        type_combo.current(0)

        # Period range
        ttk.Label(form, text="Work Period From:").grid(row=4, column=0, sticky="e", **pad)
        pfrom_frame, pfrom_var = self._create_date_entry(form, default="")
        pfrom_frame.grid(row=4, column=1, sticky="w", **pad)

        ttk.Label(form, text="Work Period To:").grid(row=5, column=0, sticky="e", **pad)
        pto_frame, pto_var = self._create_date_entry(form, default="")
        pto_frame.grid(row=5, column=1, sticky="w", **pad)

        # Notes
        ttk.Label(form, text="Notes:").grid(row=6, column=0, sticky="e", **pad)
        notes_entry = ttk.Entry(form, width=30)
        notes_entry.grid(row=6, column=1, columnspan=2, sticky="w", **pad)

        def save():
            idx = sub_combo.current()
            if idx < 0:
                messagebox.showerror("Error", "Select a subcontractor.", parent=dialog)
                return
            try:
                amount = float(amt_entry.get().replace("$", "").replace(",", ""))
            except ValueError:
                messagebox.showerror("Error", "Enter a valid amount.", parent=dialog)
                return
            if amount <= 0:
                messagebox.showerror("Error", "Amount must be greater than zero.", parent=dialog)
                return

            db.add_sub_payment(
                contractor_id=sub_ids[idx],
                amount=amount,
                payment_date=date_var.get().strip(),
                payment_type=type_var.get(),
                period_from=pfrom_var.get().strip() or None,
                period_to=pto_var.get().strip() or None,
                notes=notes_entry.get().strip(),
                user_id=self.user_id,
            )
            dialog.destroy()
            self._refresh_sub_payments()

        ttk.Button(dialog, text="Save Payment", command=save).pack(pady=15)

    def _delete_sub_payment(self):
        selected = self.sub_pay_tree.focus()
        if not selected:
            messagebox.showwarning("No Selection", "Select a payment to delete.")
            return
        values = self.sub_pay_tree.item(selected, "values")
        if messagebox.askyesno("Confirm", f"Delete payment of {values[2]} to {values[1]}?"):
            db.delete_sub_payment(int(selected))
            self._refresh_sub_payments()

    def _generate_sub_statement(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Sub Payment Statement")
        dialog.geometry("420x200")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        ttk.Label(dialog, text="Generate Subcontractor Payment Statement", font=("Segoe UI", 11, "bold")).pack(pady=(15, 10))

        form = ttk.Frame(dialog)
        form.pack(fill="x", padx=20)

        ttk.Label(form, text="From:").grid(row=0, column=0, sticky="e", padx=8, pady=8)
        from_frame, from_var = self._create_date_entry(form, default=datetime.now().strftime("%Y-01-01"))
        from_frame.grid(row=0, column=1, sticky="w", padx=8, pady=8)

        ttk.Label(form, text="To:").grid(row=0, column=2, sticky="e", padx=8, pady=8)
        to_frame, to_var = self._create_date_entry(form, default=datetime.now().strftime("%Y-%m-%d"))
        to_frame.grid(row=0, column=3, sticky="w", padx=8, pady=8)

        def generate():
            date_from = from_var.get()
            date_to = to_var.get()
            dialog.destroy()
            self._preview_sub_statement(date_from, date_to)

        ttk.Button(dialog, text="Preview", command=generate).pack(pady=15)

    def _preview_sub_statement(self, date_from, date_to):
        from export_reports import _format_date_display

        summary = db.get_sub_payment_summary(date_from, date_to, user_id=self.user_id)
        grand_total = sum(row[1] for row in summary)

        preview = tk.Toplevel(self.root)
        preview.title("Subcontractor Payment Statement")
        preview.geometry("650x500")
        self._center_dialog(preview)

        canvas = tk.Canvas(preview, bg="white")
        canvas.pack(fill="both", expand=True, padx=10, pady=(10, 5))

        company = db.load_company_info(self.user_id)
        logo_path = os.path.join(APP_DIR, "logo.png")

        y = 20
        if os.path.exists(logo_path):
            logo_img = PILImage.open(logo_path)
            logo_img.thumbnail((65, 65))
            self._sub_stmt_logo = ImageTk.PhotoImage(logo_img)
            canvas.create_image(50, y, image=self._sub_stmt_logo, anchor="nw")

        header_x = 130 if os.path.exists(logo_path) else 50
        company_name = company.get("Company Name", "").upper()
        canvas.create_text(header_x, y + 2, text=company_name, font=("Segoe UI", 16, "bold"), anchor="nw")
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company["Email"])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")
        if info_parts:
            canvas.create_text(header_x, y + 24, text="  |  ".join(info_parts), font=("Segoe UI", 9, "bold"), anchor="nw")

        y = 100
        canvas.create_line(40, y, 610, y, width=2)
        y += 15
        canvas.create_text(325, y, text="SUBCONTRACTOR PAYMENT STATEMENT", font=("Segoe UI", 13, "bold"), anchor="center")
        y += 22
        period_text = f"{_format_date_display(date_from)} – {_format_date_display(date_to)}"
        canvas.create_text(325, y, text=period_text, font=("Segoe UI", 10), anchor="center")
        y += 25
        canvas.create_line(40, y, 610, y, width=1)
        y += 15

        # Column headers
        col_name = 60
        col_count = 380
        col_total = 560

        canvas.create_text(col_name, y, text="SUBCONTRACTOR", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(col_count, y, text="# PAYMENTS", font=("Segoe UI", 9, "bold"), anchor="e")
        canvas.create_text(col_total, y, text="TOTAL PAID", font=("Segoe UI", 9, "bold"), anchor="e")
        y += 16
        canvas.create_line(40, y, 610, y, width=0.5)
        y += 10

        if not summary:
            canvas.create_text(325, y, text="No payments recorded for this period.", font=("Segoe UI", 10, "italic"), anchor="center", fill="#666666")
            y += 20
        else:
            for name, total, count in summary:
                canvas.create_text(col_name, y, text=name, font=("Segoe UI", 9), anchor="w")
                canvas.create_text(col_count, y, text=str(count), font=("Segoe UI", 9), anchor="e")
                canvas.create_text(col_total, y, text=f"${total:,.2f}", font=("Segoe UI", 9), anchor="e")
                y += 18

        y += 8
        canvas.create_line(40, y, 610, y, width=1.5)
        y += 12
        canvas.create_text(col_name, y, text="TOTAL PAID TO ALL SUBS", font=("Segoe UI", 10, "bold"), anchor="w")
        canvas.create_text(col_total, y, text=f"${grand_total:,.2f}", font=("Segoe UI", 10, "bold"), anchor="e")
        y += 20

        canvas.configure(scrollregion=(0, 0, 650, max(y + 30, 500)))

        btn_frame = ttk.Frame(preview)
        btn_frame.pack(fill="x", padx=10, pady=8)

        def export_pdf():
            filepath = filedialog.asksaveasfilename(parent=preview, title="Save as PDF", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
            if filepath:
                self._export_sub_statement_pdf(filepath, date_from, date_to, summary, grand_total)
                messagebox.showinfo("Saved", f"PDF saved to:\n{filepath}", parent=preview)

        ttk.Button(btn_frame, text="Export PDF", command=export_pdf).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Close", command=preview.destroy).pack(side="right", padx=5)

    def _export_sub_statement_pdf(self, filepath, date_from, date_to, summary, grand_total):
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from export_reports import _format_date_display

        company = db.load_company_info(self.user_id)
        doc = SimpleDocTemplate(filepath, pagesize=letter, topMargin=0.5 * inch)
        styles = getSampleStyleSheet()
        elements = []

        header_style = ParagraphStyle("H", parent=styles["Heading1"], fontSize=16, spaceAfter=4, fontName="Helvetica-Bold")
        sub_style = ParagraphStyle("S", parent=styles["Normal"], fontSize=10, spaceAfter=2, fontName="Helvetica-Bold")
        title_style = ParagraphStyle("T", parent=styles["Heading2"], fontSize=13, spaceAfter=4)
        period_style = ParagraphStyle("P", parent=styles["Normal"], fontSize=10, spaceAfter=2)

        company_name = company.get("Company Name", "").upper()
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company["Email"])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")

        logo_path = os.path.join(APP_DIR, "logo.png")
        if os.path.exists(logo_path):
            logo = Image(logo_path, width=0.9 * inch, height=0.9 * inch)
            header_content = []
            if company_name:
                header_content.append(Paragraph(company_name, header_style))
            if info_parts:
                header_content.append(Paragraph("  |  ".join(info_parts), sub_style))
            header_table_data = [[logo, header_content]]
            header_table = Table(header_table_data, colWidths=[1.1 * inch, 5.3 * inch])
            header_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))
            elements.append(header_table)
        else:
            if company_name:
                elements.append(Paragraph(company_name, header_style))
            if info_parts:
                elements.append(Paragraph("  |  ".join(info_parts), sub_style))

        elements.append(Spacer(1, 14))
        elements.append(Paragraph("Subcontractor Payment Statement", title_style))
        elements.append(Paragraph(f"{_format_date_display(date_from)} – {_format_date_display(date_to)}", period_style))
        elements.append(Spacer(1, 15))

        data = [["Subcontractor", "# Payments", "Total Paid"]]
        for name, total, count in summary:
            data.append([name, str(count), f"${total:,.2f}"])
        data.append(["TOTAL", "", f"${grand_total:,.2f}"])

        table = Table(data, colWidths=[3.5 * inch, 1.2 * inch, 1.5 * inch])
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("LINEBELOW", (0, 0), (-1, 0), 1, colors.black),
            ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ]))
        elements.append(table)
        doc.build(elements)

    def _generate_individual_sub_statement(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Individual Sub Statement")
        dialog.geometry("450x260")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        ttk.Label(dialog, text="Individual Subcontractor Statement", font=("Segoe UI", 11, "bold")).pack(pady=(15, 10))

        form = ttk.Frame(dialog)
        form.pack(fill="x", padx=20)

        ttk.Label(form, text="Subcontractor:").grid(row=0, column=0, sticky="e", padx=8, pady=8)
        subs = db.get_contractors("Subcontractor", user_id=self.user_id)
        sub_names = [c[2] for c in subs]
        sub_ids = [c[0] for c in subs]
        sub_var = tk.StringVar()
        sub_combo = ttk.Combobox(form, textvariable=sub_var, values=sub_names, width=25, state="readonly")
        sub_combo.grid(row=0, column=1, columnspan=3, sticky="w", padx=8, pady=8)

        ttk.Label(form, text="From:").grid(row=1, column=0, sticky="e", padx=8, pady=8)
        from_frame, from_var = self._create_date_entry(form, default=datetime.now().strftime("%Y-01-01"))
        from_frame.grid(row=1, column=1, sticky="w", padx=8, pady=8)

        ttk.Label(form, text="To:").grid(row=1, column=2, sticky="e", padx=8, pady=8)
        to_frame, to_var = self._create_date_entry(form, default=datetime.now().strftime("%Y-%m-%d"))
        to_frame.grid(row=1, column=3, sticky="w", padx=8, pady=8)

        def generate():
            idx = sub_combo.current()
            if idx < 0:
                messagebox.showerror("Error", "Select a subcontractor.", parent=dialog)
                return
            contractor_id = sub_ids[idx]
            contractor_name = sub_names[idx]
            date_from = from_var.get()
            date_to = to_var.get()
            dialog.destroy()
            self._preview_individual_sub_statement(contractor_id, contractor_name, date_from, date_to)

        ttk.Button(dialog, text="Preview", command=generate).pack(pady=15)

    def _preview_individual_sub_statement(self, contractor_id, contractor_name, date_from, date_to):
        from export_reports import _format_date_display

        payments = db.get_sub_payments(start_date=date_from, end_date=date_to, contractor_id=contractor_id, user_id=self.user_id)
        grand_total = sum(p[2] for p in payments)

        preview = tk.Toplevel(self.root)
        preview.title(f"Statement: {contractor_name}")
        preview.geometry("700x550")
        self._center_dialog(preview)

        canvas = tk.Canvas(preview, bg="white")
        canvas.pack(fill="both", expand=True, padx=10, pady=(10, 5))

        company = db.load_company_info(self.user_id)
        logo_path = os.path.join(APP_DIR, "logo.png")

        y = 20
        if os.path.exists(logo_path):
            logo_img = PILImage.open(logo_path)
            logo_img.thumbnail((65, 65))
            self._ind_sub_logo = ImageTk.PhotoImage(logo_img)
            canvas.create_image(50, y, image=self._ind_sub_logo, anchor="nw")

        header_x = 130 if os.path.exists(logo_path) else 50
        company_name = company.get("Company Name", "").upper()
        canvas.create_text(header_x, y + 2, text=company_name, font=("Segoe UI", 16, "bold"), anchor="nw")
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company["Email"])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")
        if info_parts:
            canvas.create_text(header_x, y + 24, text="  |  ".join(info_parts), font=("Segoe UI", 9, "bold"), anchor="nw")

        y = 100
        canvas.create_line(40, y, 660, y, width=2)
        y += 15
        canvas.create_text(350, y, text="INDIVIDUAL PAYMENT STATEMENT", font=("Segoe UI", 13, "bold"), anchor="center")
        y += 22
        canvas.create_text(350, y, text=contractor_name, font=("Segoe UI", 11, "bold"), anchor="center")
        y += 20
        period_text = f"{_format_date_display(date_from)} – {_format_date_display(date_to)}"
        canvas.create_text(350, y, text=period_text, font=("Segoe UI", 10), anchor="center")
        y += 25
        canvas.create_line(40, y, 660, y, width=1)
        y += 15

        # Column headers
        col_date = 60
        col_amount = 220
        col_type = 340
        col_period = 460
        col_notes = 620

        canvas.create_text(col_date, y, text="DATE", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(col_amount, y, text="AMOUNT", font=("Segoe UI", 9, "bold"), anchor="e")
        canvas.create_text(col_type, y, text="TYPE", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(col_period, y, text="WORK PERIOD", font=("Segoe UI", 9, "bold"), anchor="w")
        canvas.create_text(col_notes, y, text="NOTES", font=("Segoe UI", 9, "bold"), anchor="w")
        y += 16
        canvas.create_line(40, y, 660, y, width=0.5)
        y += 10

        if not payments:
            canvas.create_text(350, y, text="No payments recorded for this period.", font=("Segoe UI", 10, "italic"), anchor="center", fill="#666666")
            y += 20
        else:
            for p in payments:
                pid, name, amount, pdate, ptype, pfrom, pto, notes = p
                canvas.create_text(col_date, y, text=pdate or "", font=("Segoe UI", 9), anchor="w")
                canvas.create_text(col_amount, y, text=f"${amount:,.2f}", font=("Segoe UI", 9), anchor="e")
                canvas.create_text(col_type, y, text=ptype or "", font=("Segoe UI", 9), anchor="w")
                period = ""
                if pfrom and pto:
                    period = f"{pfrom} – {pto}"
                elif pfrom:
                    period = f"{pfrom} –"
                elif pto:
                    period = f"– {pto}"
                canvas.create_text(col_period, y, text=period, font=("Segoe UI", 9), anchor="w")
                canvas.create_text(col_notes, y, text=notes or "", font=("Segoe UI", 8), anchor="w")
                y += 18

        y += 8
        canvas.create_line(40, y, 660, y, width=1.5)
        y += 12
        canvas.create_text(col_date, y, text="TOTAL", font=("Segoe UI", 10, "bold"), anchor="w")
        canvas.create_text(col_amount, y, text=f"${grand_total:,.2f}", font=("Segoe UI", 10, "bold"), anchor="e")
        y += 18
        canvas.create_text(col_date, y, text=f"{len(payments)} payment(s)", font=("Segoe UI", 9, "italic"), anchor="w", fill="#444444")
        y += 20

        canvas.configure(scrollregion=(0, 0, 700, max(y + 30, 500)))

        btn_frame = ttk.Frame(preview)
        btn_frame.pack(fill="x", padx=10, pady=8)

        def export_pdf():
            filepath = filedialog.asksaveasfilename(parent=preview, title="Save as PDF", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
            if filepath:
                self._export_individual_sub_pdf(filepath, contractor_name, date_from, date_to, payments, grand_total)
                messagebox.showinfo("Saved", f"PDF saved to:\n{filepath}", parent=preview)

        ttk.Button(btn_frame, text="Export PDF", command=export_pdf).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Close", command=preview.destroy).pack(side="right", padx=5)

    def _export_individual_sub_pdf(self, filepath, contractor_name, date_from, date_to, payments, grand_total):
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from export_reports import _format_date_display

        company = db.load_company_info(self.user_id)
        doc = SimpleDocTemplate(filepath, pagesize=letter, topMargin=0.5 * inch)
        styles = getSampleStyleSheet()
        elements = []

        header_style = ParagraphStyle("H", parent=styles["Heading1"], fontSize=16, spaceAfter=4, fontName="Helvetica-Bold")
        sub_style = ParagraphStyle("S", parent=styles["Normal"], fontSize=10, spaceAfter=2, fontName="Helvetica-Bold")
        title_style = ParagraphStyle("T", parent=styles["Heading2"], fontSize=13, spaceAfter=4)
        period_style = ParagraphStyle("P", parent=styles["Normal"], fontSize=10, spaceAfter=2)

        company_name = company.get("Company Name", "").upper()
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company["Email"])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")

        logo_path = os.path.join(APP_DIR, "logo.png")
        if os.path.exists(logo_path):
            logo = Image(logo_path, width=0.9 * inch, height=0.9 * inch)
            header_content = []
            if company_name:
                header_content.append(Paragraph(company_name, header_style))
            if info_parts:
                header_content.append(Paragraph("  |  ".join(info_parts), sub_style))
            header_table_data = [[logo, header_content]]
            header_table = Table(header_table_data, colWidths=[1.1 * inch, 5.3 * inch])
            header_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))
            elements.append(header_table)
        else:
            if company_name:
                elements.append(Paragraph(company_name, header_style))
            if info_parts:
                elements.append(Paragraph("  |  ".join(info_parts), sub_style))

        elements.append(Spacer(1, 14))
        elements.append(Paragraph("Individual Payment Statement", title_style))
        elements.append(Paragraph(f"<b>{contractor_name}</b>", period_style))
        elements.append(Paragraph(f"{_format_date_display(date_from)} – {_format_date_display(date_to)}", period_style))
        elements.append(Spacer(1, 15))

        data = [["Date", "Amount", "Type", "Work Period", "Notes"]]
        for p in payments:
            pid, name, amount, pdate, ptype, pfrom, pto, notes = p
            period = ""
            if pfrom and pto:
                period = f"{pfrom} – {pto}"
            elif pfrom:
                period = f"{pfrom} –"
            elif pto:
                period = f"– {pto}"
            data.append([pdate or "", f"${amount:,.2f}", ptype or "", period, notes or ""])
        data.append(["TOTAL", f"${grand_total:,.2f}", "", f"{len(payments)} payment(s)", ""])

        table = Table(data, colWidths=[0.9 * inch, 1 * inch, 0.7 * inch, 1.8 * inch, 2 * inch])
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("LINEBELOW", (0, 0), (-1, 0), 1, colors.black),
            ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ]))
        elements.append(table)
        doc.build(elements)

    # ──────────────────────────────────────────────
    # Reports
    # ──────────────────────────────────────────────
    def _generate_pl(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Generate Profit & Loss")
        dialog.geometry("450x240")
        dialog.transient(self.root)
        dialog.grab_set()
        self._center_dialog(dialog)
        dialog.resizable(False, False)

        ttk.Label(dialog, text="Profit & Loss Statement", font=("Segoe UI", 12, "bold")).pack(pady=(15, 10))

        form = ttk.Frame(dialog)
        form.pack(fill="x", padx=20)

        ttk.Label(form, text="From:").grid(row=0, column=0, sticky="e", padx=8, pady=8)
        from_frame, from_var = self._create_date_entry(form, default=datetime.now().strftime("%Y-01-01"))
        from_frame.grid(row=0, column=1, sticky="w", padx=8, pady=8)

        ttk.Label(form, text="To:").grid(row=0, column=2, sticky="e", padx=8, pady=8)
        to_frame, to_var = self._create_date_entry(form, default=datetime.now().strftime("%Y-%m-%d"))
        to_frame.grid(row=0, column=3, sticky="w", padx=8, pady=8)

        compare_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(form, text="Include prior period comparison", variable=compare_var).grid(row=1, column=0, columnspan=4, sticky="w", padx=8, pady=8)

        def generate():
            date_from = from_var.get()
            date_to = to_var.get()
            dialog.destroy()
            self._preview_pl(date_from, date_to, compare_var.get())

        ttk.Button(dialog, text="Preview", command=generate).pack(pady=15)

    def _preview_pl(self, date_from, date_to, compare):
        from export_reports import _compute_pl, _prior_period, _format_date_display

        transactions = db.get_transactions(date_from, date_to, user_id=self.user_id)
        income, cogs, expenses = _compute_pl(transactions)

        total_income = sum(income.values())
        total_cogs = sum(cogs.values())
        gross_profit = total_income - total_cogs
        total_expenses = sum(expenses.values())
        net_income = gross_profit - total_expenses
        margin = (net_income / total_income * 100) if total_income != 0 else 0

        prior_income = {}
        prior_cogs = {}
        prior_expenses = {}
        if compare:
            prior_from, prior_to = _prior_period(date_from, date_to)
            prior_txns = db.get_transactions(prior_from, prior_to, user_id=self.user_id)
            prior_income, prior_cogs, prior_expenses = _compute_pl(prior_txns)

        preview = tk.Toplevel(self.root)
        preview.title("Profit & Loss Preview")
        preview.geometry("750x600")
        self._center_dialog(preview)

        canvas = tk.Canvas(preview, bg="white")
        canvas.pack(fill="both", expand=True, padx=10, pady=(10, 5))

        company = db.load_company_info(self.user_id)
        logo_path = os.path.join(APP_DIR, "logo.png")

        y = 20
        if os.path.exists(logo_path):
            logo_img = PILImage.open(logo_path)
            logo_img.thumbnail((65, 65))
            self._pl_preview_logo = ImageTk.PhotoImage(logo_img)
            canvas.create_image(50, y, image=self._pl_preview_logo, anchor="nw")

        header_x = 130 if os.path.exists(logo_path) else 50
        company_name = company.get("Company Name", "").upper()
        canvas.create_text(header_x, y + 2, text=company_name, font=("Segoe UI", 16, "bold"), anchor="nw")
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company["Email"])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")
        if info_parts:
            canvas.create_text(header_x, y + 24, text="  |  ".join(info_parts), font=("Segoe UI", 9, "bold"), anchor="nw")

        y = 100
        canvas.create_line(40, y, 710, y, width=2)
        y += 15
        canvas.create_text(375, y, text="PROFIT & LOSS STATEMENT", font=("Segoe UI", 13, "bold"), anchor="center")
        y += 22
        period_text = f"{_format_date_display(date_from)} – {_format_date_display(date_to)}"
        canvas.create_text(375, y, text=period_text, font=("Segoe UI", 10), anchor="center")
        y += 25
        canvas.create_line(40, y, 710, y, width=1)
        y += 15

        col_cat = 60
        col_cur = 500
        col_prior = 600
        col_chg = 690

        if compare:
            canvas.create_text(col_cat, y, text="CATEGORY", font=("Segoe UI", 9, "bold"), anchor="w")
            canvas.create_text(col_cur, y, text="CURRENT", font=("Segoe UI", 9, "bold"), anchor="e")
            canvas.create_text(col_prior, y, text="PRIOR", font=("Segoe UI", 9, "bold"), anchor="e")
            canvas.create_text(col_chg, y, text="CHANGE", font=("Segoe UI", 9, "bold"), anchor="e")
        else:
            canvas.create_text(col_cat, y, text="CATEGORY", font=("Segoe UI", 9, "bold"), anchor="w")
            canvas.create_text(col_chg, y, text="AMOUNT", font=("Segoe UI", 9, "bold"), anchor="e")
        y += 16
        canvas.create_line(40, y, 710, y, width=0.5)
        y += 10

        def draw_row(label, current, prior=0, bold=False):
            nonlocal y
            font = ("Segoe UI", 9, "bold") if bold else ("Segoe UI", 9)
            canvas.create_text(col_cat, y, text=label, font=font, anchor="w")
            canvas.create_text(col_cur if compare else col_chg, y, text=f"${current:,.2f}", font=font, anchor="e")
            if compare:
                canvas.create_text(col_prior, y, text=f"${prior:,.2f}", font=font, anchor="e")
                change = current - prior
                canvas.create_text(col_chg, y, text=f"${change:+,.2f}", font=font, anchor="e")
            y += 16

        def draw_section_header(label):
            nonlocal y
            canvas.create_text(col_cat, y, text=label, font=("Segoe UI", 9, "bold"), anchor="w")
            y += 16

        draw_section_header("INCOME")
        all_income = sorted(set(list(income.keys()) + list(prior_income.keys())))
        for name in all_income:
            draw_row(f"  {name}", income.get(name, 0), prior_income.get(name, 0))
        draw_row("Total Income", total_income, sum(prior_income.values()), bold=True)
        y += 8

        draw_section_header("COST OF GOODS SOLD")
        all_cogs = sorted(set(list(cogs.keys()) + list(prior_cogs.keys())))
        for name in all_cogs:
            draw_row(f"  {name}", cogs.get(name, 0), prior_cogs.get(name, 0))
        draw_row("Total COGS", total_cogs, sum(prior_cogs.values()), bold=True)
        y += 8

        prior_gross = sum(prior_income.values()) - sum(prior_cogs.values())
        draw_row("GROSS PROFIT", gross_profit, prior_gross, bold=True)
        y += 8

        draw_section_header("EXPENSES")
        all_exp = sorted(set(list(expenses.keys()) + list(prior_expenses.keys())))
        for name in all_exp:
            draw_row(f"  {name}", expenses.get(name, 0), prior_expenses.get(name, 0))
        draw_row("Total Expenses", total_expenses, sum(prior_expenses.values()), bold=True)
        y += 8

        canvas.create_line(40, y, 710, y, width=1.5)
        y += 12
        prior_net = prior_gross - sum(prior_expenses.values())
        draw_row("NET INCOME", net_income, prior_net, bold=True)
        y += 5
        canvas.create_text(col_cat, y, text=f"Net Profit Margin: {margin:.1f}%", font=("Segoe UI", 9, "italic"), anchor="w", fill="#444444")
        y += 20

        canvas.configure(scrollregion=(0, 0, 750, max(y + 30, 600)))

        btn_frame = ttk.Frame(preview)
        btn_frame.pack(fill="x", padx=10, pady=8)

        def export_pdf():
            filepath = filedialog.asksaveasfilename(parent=preview, title="Save as PDF", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
            if filepath:
                export_reports.export_pl_pdf(filepath, date_from, date_to, compare=compare, user_id=self.user_id)
                messagebox.showinfo("Saved", f"PDF saved to:\n{filepath}", parent=preview)

        def export_csv():
            filepath = filedialog.asksaveasfilename(parent=preview, title="Save as CSV", defaultextension=".csv", filetypes=[("CSV", "*.csv")])
            if filepath:
                export_reports.export_pl_csv(filepath, date_from, date_to, compare=compare, user_id=self.user_id)
                messagebox.showinfo("Saved", f"CSV saved to:\n{filepath}", parent=preview)

        ttk.Button(btn_frame, text="Export PDF", command=export_pdf).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Export CSV", command=export_csv).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Close", command=preview.destroy).pack(side="right", padx=5)

    def _generate_balance_sheet(self):
        self._preview_balance_sheet()

    def _preview_balance_sheet(self):
        transactions = db.get_transactions(user_id=self.user_id)
        from export_reports import _compute_bs
        assets, liabilities, equity = _compute_bs(transactions)

        total_assets = sum(assets.values())
        total_liabilities = sum(liabilities.values())
        total_equity = sum(equity.values())

        preview = tk.Toplevel(self.root)
        preview.title("Balance Sheet Preview")
        preview.geometry("650x550")
        self._center_dialog(preview)

        canvas = tk.Canvas(preview, bg="white")
        canvas.pack(fill="both", expand=True, padx=10, pady=(10, 5))

        company = db.load_company_info(self.user_id)
        logo_path = os.path.join(APP_DIR, "logo.png")

        y = 20
        if os.path.exists(logo_path):
            logo_img = PILImage.open(logo_path)
            logo_img.thumbnail((65, 65))
            self._bs_preview_logo = ImageTk.PhotoImage(logo_img)
            canvas.create_image(50, y, image=self._bs_preview_logo, anchor="nw")

        header_x = 130 if os.path.exists(logo_path) else 50
        company_name = company.get("Company Name", "").upper()
        canvas.create_text(header_x, y + 2, text=company_name, font=("Segoe UI", 16, "bold"), anchor="nw")
        phone_fmt = self._format_phone_display(company.get("Phone Number", ""))
        info_parts = []
        if phone_fmt:
            info_parts.append(phone_fmt)
        if company.get("Email"):
            info_parts.append(company["Email"])
        if company.get("HIC Number"):
            info_parts.append(f"HIC# {company['HIC Number']}")
        if info_parts:
            canvas.create_text(header_x, y + 24, text="  |  ".join(info_parts), font=("Segoe UI", 9, "bold"), anchor="nw")

        y = 100
        canvas.create_line(40, y, 610, y, width=2)
        y += 15
        canvas.create_text(325, y, text="BALANCE SHEET", font=("Segoe UI", 13, "bold"), anchor="center")
        y += 22
        canvas.create_text(325, y, text=f"As of {datetime.now().strftime('%B %d, %Y')}", font=("Segoe UI", 10), anchor="center")
        y += 25
        canvas.create_line(40, y, 610, y, width=1)
        y += 15

        col_cat = 60
        col_amt = 560

        def draw_row(label, amount, bold=False):
            nonlocal y
            font = ("Segoe UI", 9, "bold") if bold else ("Segoe UI", 9)
            canvas.create_text(col_cat, y, text=label, font=font, anchor="w")
            canvas.create_text(col_amt, y, text=f"${amount:,.2f}", font=font, anchor="e")
            y += 16

        def draw_section_header(label):
            nonlocal y
            canvas.create_text(col_cat, y, text=label, font=("Segoe UI", 9, "bold"), anchor="w")
            y += 16

        draw_section_header("ASSETS")
        for name, val in sorted(assets.items()):
            draw_row(f"  {name}", val)
        draw_row("Total Assets", total_assets, bold=True)
        y += 10

        draw_section_header("LIABILITIES")
        for name, val in sorted(liabilities.items()):
            draw_row(f"  {name}", val)
        draw_row("Total Liabilities", total_liabilities, bold=True)
        y += 10

        draw_section_header("EQUITY")
        for name, val in sorted(equity.items()):
            draw_row(f"  {name}", val)
        draw_row("Total Equity", total_equity, bold=True)
        y += 10

        canvas.create_line(40, y, 610, y, width=1.5)
        y += 12
        draw_row("LIABILITIES + EQUITY", total_liabilities + total_equity, bold=True)

        canvas.configure(scrollregion=(0, 0, 650, max(y + 30, 550)))

        btn_frame = ttk.Frame(preview)
        btn_frame.pack(fill="x", padx=10, pady=8)

        def export_pdf():
            filepath = filedialog.asksaveasfilename(parent=preview, title="Save as PDF", defaultextension=".pdf", filetypes=[("PDF", "*.pdf")])
            if filepath:
                export_reports.export_bs_pdf(filepath, user_id=self.user_id)
                messagebox.showinfo("Saved", f"PDF saved to:\n{filepath}", parent=preview)

        def export_csv():
            filepath = filedialog.asksaveasfilename(parent=preview, title="Save as CSV", defaultextension=".csv", filetypes=[("CSV", "*.csv")])
            if filepath:
                export_reports.export_bs_csv(filepath, user_id=self.user_id)
                messagebox.showinfo("Saved", f"CSV saved to:\n{filepath}", parent=preview)

        ttk.Button(btn_frame, text="Export PDF", command=export_pdf).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Export CSV", command=export_csv).pack(side="left", padx=5)
        ttk.Button(btn_frame, text="Close", command=preview.destroy).pack(side="right", padx=5)


if __name__ == "__main__":
    db.backup_database()
    db.init_db()

    login = auth_ui.LoginWindow()
    if login.user_id is None:
        sys.exit(0)

    root = tk.Tk()
    app = BookkeepingApp(root, login.user_id)
    root.mainloop()
